from __future__ import annotations

import csv
import email.utils
import smtplib
import ssl
from email.message import EmailMessage
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError
import io
import json
import os
import platform
import socket
import re
import shutil
import sys
import subprocess
import threading
import time
import urllib.request
import xml.etree.ElementTree as ET
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import parse_qs, urlencode, urlsplit, quote

from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, BarChart, PieChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
VENDOR = ROOT / "vendor"
if VENDOR.exists(): sys.path.insert(0, str(VENDOR))
try:
    import MetaTrader5 as mt5
except ImportError:
    mt5 = None
JOURNAL = Path(os.environ.get("MT5_JOURNAL_PATH", ROOT / "TradeJournal.csv"))
BOOK = ROOT / "Xuzera-TradingOS.xlsx"
SCREENSHOTS = ROOT / "Screenshots"
HOST, PORT = "127.0.0.1", int(os.environ.get("PORT", "8766"))
OLLAMA_URL = os.environ.get("OLLAMA_URL", "http://127.0.0.1:11434").rstrip("/")
MODEL_CONFIG = ROOT / "settings.json"
OLLAMA_MODEL = os.environ.get("OLLAMA_MODEL", "qwen3:14b")
try:
    OLLAMA_MODEL = json.loads(MODEL_CONFIG.read_text(encoding="utf-8")).get("model", OLLAMA_MODEL)
except (OSError, ValueError, AttributeError):
    pass
OLLAMA_LOCK = threading.RLock()
OLLAMA_LAST_START = 0.0
OLLAMA_PULL = {"running": False, "model": "", "status": "", "completed": 0, "total": 0, "error": ""}
OLLAMA_RECOMMENDED = [
    {"name": "qwen3.5:4b", "label": "Qwen 3.5 · 4B", "description": "Smaller download; a good fit for lighter computers."},
    {"name": "qwen3.5:9b", "label": "Qwen 3.5 · 9B", "description": "Balanced general analysis; requires more memory."},
    {"name": "qwen3:14b", "label": "Qwen 3 · 14B", "description": "Stronger analysis; larger download and memory use."},
    {"name": "llama3.1:8b", "label": "Llama 3.1 · 8B", "description": "Alternative general-purpose model."},
    {"name": "deepseek-r1:14b", "label": "DeepSeek R1 · 14B", "description": "Reasoning-focused; can respond more slowly."},
    {"name": "gemma4:26b", "label": "Gemma 4 · 26B", "description": "Larger general model; needs substantial memory."},
]
DASHBOARD_FILE = ROOT / "dashboard.html"
APP_ICON = ROOT / "Xuzera.ico"
MT5_EXE = os.environ.get("MT5_PATH", r"C:\Program Files\MetaTrader 5\terminal64.exe")
COMMON_FILES = Path(os.environ.get("APPDATA", str(ROOT))) / "MetaQuotes" / "Terminal" / "Common" / "Files"
EA_HEARTBEAT = COMMON_FILES / "GoldTradingControlCenter_EA.txt"
MT5_DATA_DIR = Path(os.environ.get("MT5_DATA_PATH", str(Path(os.environ.get("APPDATA", str(ROOT))) / "MetaQuotes" / "Terminal" / "D0E8209F77C8CF37AD8BF550E51FF075")))
MT5_EXPERT_LOG_DIR = MT5_DATA_DIR / "MQL5" / "Logs"
MT5_LOCK = threading.RLock()
MT5_CACHE = {"at": 0.0, "data": None, "initialized": False}
MT5_CACHE_TTL = 1.5
MT5_HISTORY_CACHE = {"at": 0.0, "rows": [], "error": "", "available": False, "sampled_at": ""}
MT5_HISTORY_TTL = 1
NEWS_URL = "https://api.gdeltproject.org/api/v2/doc/doc"
NEWS_TTL = 300
NEWS_LOCK = threading.Lock()
NEWS_CACHE = {"at": 0.0, "items": [], "updated_at": "", "error": "", "loading": False, "provider": ""}
NEWS_QUERY = '(gold OR bullion OR "Federal Reserve" OR inflation OR "US dollar" OR yields OR tariffs OR Iran OR Israel OR Ukraine)'
NEWS_REFRESH_LOCK = threading.Lock()
NEWS_REFRESH_RUNNING = False
MARKET_PULSE_API_BASE = os.environ.get("MARKET_PULSE_API_BASE", "http://127.0.0.1:8000/api/v1").strip().rstrip("/")
MARKET_PULSE_TIMEOUT = 5
MARKET_PULSE_STATUS = {"configured": bool(MARKET_PULSE_API_BASE), "available": False, "count": 0, "checked_at": "", "error": "Not checked yet"}
FOREX_FACTORY_CALENDAR_URL = "https://nfs.faireconomy.media/ff_calendar_thisweek.json"
CALENDAR_TTL = 600
CALENDAR_LOCK = threading.Lock()
CALENDAR_REFRESH_LOCK = threading.Lock()
CALENDAR_REFRESH_RUNNING = False
CALENDAR_CACHE = {"at": 0.0, "events": [], "updated_at": "", "error": "", "loading": False, "provider": "Forex Factory"}
FIELDS = ["Trade ID", "Ticket Number", "Date", "Open Time", "Close Time", "Trade Duration", "Trade Status", "Symbol", "Trade Type", "Lot Size", "Entry Price", "Stop Loss", "Take Profit", "Exit Price", "Profit/Loss", "Account Currency", "Spread", "Risk %", "Win/Loss", "Strategy Name", "Signal Type", "Entry Reason", "Exit Reason", "Exit Comment", "Screenshot", "R Multiple", "Balance", "Equity"]

PAGE = r'''<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Xuzera</title><style>
:root{color-scheme:dark;--bg:#0b1020;--panel:#121a2d;--line:#26324b;--muted:#9aa9c3;--gold:#f2bd56;--green:#4ed6a0;--red:#ff7777}*{box-sizing:border-box}body{margin:0;background:var(--bg);color:#edf2fb;font:14px system-ui,Segoe UI,sans-serif}header{padding:25px 4vw;border-bottom:1px solid var(--line);display:flex;justify-content:space-between;align-items:center}h1{font-size:22px;margin:0}small,.muted{color:var(--muted)}button{border:1px solid #b98736;background:#201b12;color:var(--gold);padding:10px 14px;border-radius:8px;cursor:pointer;font-weight:650}main{padding:24px 4vw;max-width:1500px;margin:auto}.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(165px,1fr));gap:12px}.card,.panel{background:var(--panel);border:1px solid var(--line);border-radius:12px;padding:17px}.label{color:var(--muted);font-size:12px}.value{font-size:24px;font-weight:700;margin-top:8px}.row{display:grid;grid-template-columns:1.35fr 1fr;gap:14px;margin-top:14px}.panel{overflow:auto}.panel h2{font-size:15px;margin:0 0 14px}table{border-collapse:collapse;width:100%;white-space:nowrap}th,td{text-align:left;padding:10px 9px;border-bottom:1px solid var(--line)}th{color:var(--muted);font-size:11px;text-transform:uppercase}td{font-size:12px}.positive{color:var(--green)}.negative{color:var(--red)}.notice{margin-top:14px;padding:12px 15px;border:1px solid #594827;border-radius:10px;color:#e6c989;background:#211c13}.charts{display:grid;grid-template-columns:1fr 1fr;gap:16px}.bar{height:180px;display:flex;gap:5px;align-items:end;border-bottom:1px solid var(--line);padding-top:20px}.bar span{flex:1;background:linear-gradient(#eabd63,#9b6923);border-radius:4px 4px 0 0;min-width:3px;position:relative}.bar span i{position:absolute;bottom:-21px;left:0;font-style:normal;font-size:9px;color:var(--muted)}#analysis{white-space:pre-wrap;line-height:1.6;color:#d5deed}.footer{margin:20px 0;color:var(--muted);font-size:12px}@media(max-width:850px){.row,.charts{grid-template-columns:1fr}header{gap:12px;align-items:flex-start}}
</style></head><body><header><div><h1>Xuzera</h1><small>MT5 journal · local analytics · AI is optional and analysis-only</small></div><button onclick="analyze()">Generate AI review</button></header><main><div class="grid" id="metrics"></div><div class="notice">Trading is controlled by the MT5 Expert Advisor, not this dashboard. The EA ships with trading disabled. Review on a demo account before enabling live orders.</div><div class="row"><section class="panel"><h2>Equity and profit curve</h2><div class="charts"><div id="equity"></div><div id="months"></div></div></section><section class="panel"><h2>Performance snapshot</h2><div id="summary"></div></section></div><div class="row"><section class="panel"><h2>Recent trades</h2><div style="overflow:auto"><table><thead><tr><th>Closed</th><th>Ticket</th><th>Side</th><th>Lots</th><th>Entry</th><th>Exit</th><th>P/L</th><th>Exit reason</th></tr></thead><tbody id="trades"></tbody></table></div></section><section class="panel"><h2>AI performance review</h2><div id="analysis" class="muted">Click “Generate AI review” to ask your local Ollama model. Trading and journaling continue if Ollama is offline.</div></section></div><p class="footer" id="stamp"></p></main><script>
const money=x=>Number(x||0).toLocaleString(undefined,{style:'currency',currency:'USD',maximumFractionDigits:2});function bars(el,items){if(!items.length){el.innerHTML='<small class="muted">No closed trades yet</small>';return}let max=Math.max(1,...items.map(x=>Math.abs(x.v)));el.innerHTML='<div class="bar">'+items.map(x=>`<span style="height:${Math.max(3,Math.abs(x.v)/max*100)}%" title="${x.k}: ${money(x.v)}"><i>${x.k}</i></span>`).join('')+'</div>'}async function load(){let d=await(await fetch('/api/data')).json();let m=d.metrics;let cards=[['Balance',money(m.balance)],['Equity',money(m.equity)],['Open trades',m.open_trades],['Win rate',m.win_rate.toFixed(1)+'%'],['Today P/L',money(m.daily_pl)],['Week P/L',money(m.weekly_pl)],['Month P/L',money(m.monthly_pl)],['Current drawdown',m.drawdown.toFixed(2)+'%'],['Profit factor',m.profit_factor.toFixed(2)],['Best hour',m.best_hour],['Worst hour',m.worst_hour]];document.querySelector('#metrics').innerHTML=cards.map(c=>`<div class="card"><div class="label">${c[0]}</div><div class="value">${c[1]}</div></div>`).join('');bars(document.querySelector('#equity'),d.equity.map((v,i)=>({k:i+1,v})));bars(document.querySelector('#months'),d.months.map(x=>({k:x.month.slice(5),v:x.profit})));document.querySelector('#summary').innerHTML=`<p><b>${m.total_trades}</b> closed trades · <b>${m.wins}</b> wins · <b>${m.losses}</b> losses</p><p>Average R: <b>${m.average_rr.toFixed(2)}R</b></p><p>Largest peak-to-trough drawdown: <b>${m.max_drawdown.toFixed(2)}%</b></p><p>Best hour: <b>${m.best_hour}</b> · Worst hour: <b>${m.worst_hour}</b></p>`;document.querySelector('#trades').innerHTML=d.trades.slice(-30).reverse().map(t=>`<tr><td>${t['Close Time']||''}</td><td>${t['Ticket Number']||''}</td><td>${t['Trade Type']||''}</td><td>${t['Lot Size']||''}</td><td>${t['Entry Price']||''}</td><td>${t['Exit Price']||''}</td><td class="${Number(t['Profit/Loss'])>=0?'positive':'negative'}">${money(t['Profit/Loss'])}</td><td>${t['Exit Reason']||''}</td></tr>`).join('');document.querySelector('#stamp').textContent='Updated '+new Date().toLocaleString()+' · Journal: '+d.journal}async function analyze(){let a=document.querySelector('#analysis');a.textContent='Asking the local Ollama model…';try{let r=await fetch('/api/analyze',{method:'POST'}),j=await r.json();a.textContent=j.report||j.error}catch(e){a.textContent='Could not reach the local analysis service.'}}load();setInterval(load,10000);
</script></body></html>'''


def number(value, default=0.0):
    try:
        return float(str(value).replace(",", "")) if value not in (None, "") else default
    except (ValueError, TypeError):
        return default


def excel_value(value):
    """Return a value openpyxl can safely store in an .xlsx cell."""
    if value is None or isinstance(value, (bool, int, float, date, datetime)):
        if isinstance(value, float) and not __import__("math").isfinite(value):
            return None
        return value
    if isinstance(value, (dict, list, tuple, set)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, str):
        # XML 1.0 disallows most C0 controls; openpyxl rejects them in strings.
        value = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F\uFFFE\uFFFF]", "", value)
        return value[:32767]
    return str(value)[:32767]


def excel_row(values):
    return [excel_value(value) for value in values]


def classify_gold_news(title):
    """Explain a possible gold-market channel from headline text only."""
    text = str(title or "").casefold()
    groups = [
        ("Central banks & rates", ["federal reserve", "fed ", "fomc", "interest rate", "rate cut", "rate hike", "hawkish", "dovish", "treasury yield", "bond yield"], 3),
        ("US dollar", ["us dollar", "dollar", "dxy", "greenback"], 2),
        ("Inflation & data", ["inflation", "cpi", "ppi", "jobs report", "payroll", "unemployment", "pce"], 2),
        ("Geopolitics & safe haven", ["war", "missile", "attack", "airstrike", "sanction", "invasion", "conflict", "ceasefire", "peace deal", "geopolitical"], 2),
        ("Gold demand & supply", ["gold demand", "central bank buying", "gold reserve", "bullion", "gold price", "gold market", "gold mine", "gold production"], 2),
        ("Trade & policy", ["tariff", "trade war", "import duty", "export ban"], 1),
    ]
    found = [(name, terms, weight) for name, terms, weight in groups if any(term in text for term in terms)]
    category = found[0][0] if found else "Macro & markets"
    score = 0
    reasons = []
    # Only explicit directional language is scored. Ambiguous macro headlines stay mixed.
    bullish = ["rate cut", "cuts rates", "dovish", "weaker dollar", "dollar falls", "dollar slides", "yields fall", "yields drop", "war", "missile", "attack", "airstrike", "sanction", "invasion", "escalat", "central bank buying", "gold demand rises", "gold demand surges"]
    bearish = ["rate hike", "hikes rates", "hawkish", "stronger dollar", "dollar rises", "dollar rallies", "yields rise", "yields climb", "ceasefire", "peace deal", "de-escalat", "gold demand falls"]
    if any(term in text for term in bullish):
        score += 1
        reasons.append("safe-haven demand, easier policy, or softer-dollar language can support gold")
    if any(term in text for term in bearish):
        score -= 1
        reasons.append("higher-rate, stronger-dollar, or easing-risk language can pressure gold")
    if category == "Inflation & data" and not reasons:
        reasons.append("inflation and economic data can move both rate expectations and safe-haven demand")
    elif not reasons:
        reasons.append("headline context alone does not establish a clear direction for gold")
    if score > 0:
        tilt = "Potentially supportive"
    elif score < 0:
        tilt = "Potentially negative"
    else:
        tilt = "Mixed / unclear"
    return {"category": category, "tilt": tilt, "reason": "; ".join(reasons), "confidence": "Low"}


def _news_item(title, link, source, published="", language=""):
    title, link = str(title or "").strip(), str(link or "").strip()
    parsed = urlsplit(link)
    if not title or parsed.scheme not in ("http", "https") or not parsed.netloc:
        return None
    effect = classify_gold_news(title)
    return {"title": title[:500], "url": link, "source": str(source or "Source unavailable")[:160],
            "published_at": published, "language": str(language or "")[:40], **effect}


def _fetch_google_news():
    query = 'gold OR bullion OR "Federal Reserve" OR inflation OR "US dollar" OR yields OR tariffs OR Iran OR Israel OR Ukraine when:1d'
    params = urlencode({"q": query, "hl": "en-US", "gl": "US", "ceid": "US:en"})
    req = urllib.request.Request(f"https://news.google.com/rss/search?{params}", headers={"User-Agent": "Mozilla/5.0 GoldTradingControlCenter/1.0", "Accept": "application/rss+xml, application/xml"})
    with urllib.request.urlopen(req, timeout=8) as response:
        root = ET.fromstring(response.read(1_000_000))
    items, seen = [], set()
    for node in root.findall("./channel/item"):
        source = node.findtext("source") or "Google News source"
        title = node.findtext("title") or ""
        suffix = " - " + source
        if title.endswith(suffix):
            title = title[:-len(suffix)]
        published = ""
        try:
            published = email.utils.parsedate_to_datetime(node.findtext("pubDate") or "").astimezone(timezone.utc).isoformat()
        except (TypeError, ValueError, OverflowError):
            pass
        item = _news_item(title, node.findtext("link"), source, published, "en")
        if item and item["url"] not in seen:
            seen.add(item["url"]); items.append(item)
    if not items:
        raise ValueError("Google News returned no matching headlines")
    return items[:40], "Google News RSS"


def _fetch_gdelt_news():
    params = urlencode({"query": NEWS_QUERY, "mode": "artlist", "maxrecords": "15", "timespan": "1day", "sort": "datedesc", "format": "json"})
    req = urllib.request.Request(f"{NEWS_URL}?{params}", headers={"User-Agent": "Xuzera-TradingOS/1.0 (local dashboard)", "Accept": "application/json"})
    with urllib.request.urlopen(req, timeout=8) as response:
        payload = json.loads(response.read(1_000_000).decode("utf-8"))
    items, seen = [], set()
    for article in payload.get("articles", []):
        date_raw = str(article.get("seendate") or "")
        try:
            published = datetime.strptime(date_raw[:15], "%Y%m%dT%H%M%S").replace(tzinfo=timezone.utc).isoformat()
        except ValueError:
            published = ""
        item = _news_item(article.get("title"), article.get("url"), article.get("domain"), published, article.get("language"))
        if item and item["url"] not in seen:
            seen.add(item["url"]); items.append(item)
    if not items:
        raise ValueError("GDELT returned no matching headlines")
    return items[:15], "GDELT DOC 2.0"


def _fetch_market_pulse_news():
    """Read normalized articles from a running Market Pulse API instance."""
    if not MARKET_PULSE_API_BASE:
        raise RuntimeError("Market Pulse API base URL is not configured")
    endpoint = MARKET_PULSE_API_BASE + "/news?limit=40"
    request = urllib.request.Request(endpoint, headers={
        "User-Agent": "Xuzera-TradingOS/1.0 (Market Pulse connector)",
        "Accept": "application/json",
    })
    try:
        with urllib.request.urlopen(request, timeout=MARKET_PULSE_TIMEOUT) as response:
            payload = json.loads(response.read(3_000_000).decode("utf-8"))
    except Exception as exc:
        MARKET_PULSE_STATUS.update(available=False, count=0, checked_at=datetime.now().astimezone().isoformat(timespec="seconds"), error=f"Backend unavailable ({type(exc).__name__}); using fallback news")
        raise RuntimeError("Market Pulse backend unavailable") from exc
    if not isinstance(payload, list):
        MARKET_PULSE_STATUS.update(available=False, count=0, checked_at=datetime.now().astimezone().isoformat(timespec="seconds"), error="Backend returned an unexpected response")
        raise ValueError("Market Pulse API returned an unexpected response")
    items, seen = [], set()
    for article in payload:
        if not isinstance(article, dict):
            continue
        base = _news_item(article.get("title"), article.get("url"), article.get("source"), article.get("published_at"), "en")
        if not base or base["url"] in seen:
            continue
        seen.add(base["url"])
        label = str(article.get("sentiment_label") or "neutral").casefold()
        base["tilt"] = "Potentially supportive" if label == "bullish" else "Potentially negative" if label == "bearish" else "Mixed / unclear"
        score = number(article.get("sentiment_score"), 0.0)
        context = str(article.get("overview") or article.get("excerpt") or "").strip()
        base["reason"] = (f"Market Pulse sentiment: {label} ({score:+.2f}). " + context)[:650] if context else f"Market Pulse sentiment: {label} ({score:+.2f}); verify the article and market response."
        base["confidence"] = "Provider label · verify"
        base["category"] = str(article.get("category") or base["category"])[:80]
        base["tags"] = [str(x)[:60] for x in (article.get("tags") or [])[:12] if x]
        base["related_tickers"] = [str(x)[:24] for x in (article.get("related_tickers") or [])[:20] if x]
        base["overview"] = context[:1200]
        base["key_takeaways"] = [str(x)[:240] for x in (article.get("key_takeaways") or [])[:8] if x]
        items.append(base)
    if not items:
        MARKET_PULSE_STATUS.update(available=False, count=0, checked_at=datetime.now().astimezone().isoformat(timespec="seconds"), error="Backend is reachable but returned no valid articles")
        raise ValueError("Market Pulse returned no valid news articles")
    items.sort(key=lambda item: item.get("published_at") or "", reverse=True)
    MARKET_PULSE_STATUS.update(available=True, count=len(items), checked_at=datetime.now().astimezone().isoformat(timespec="seconds"), error="")
    return items[:40], "Market Pulse"


def _news_refresh_worker():
    global NEWS_REFRESH_RUNNING
    errors = []
    try:
        items, provider_name = _fetch_market_pulse_news()
        with NEWS_LOCK:
            NEWS_CACHE.update(at=time.monotonic(), items=items, updated_at=datetime.now().astimezone().isoformat(timespec="seconds"), error="", loading=False, provider=provider_name)
    except Exception as exc:
        errors.append(str(exc))
        succeeded = False
        for provider in (_fetch_google_news, _fetch_gdelt_news):
            try:
                items, provider_name = provider()
                with NEWS_LOCK:
                    NEWS_CACHE.update(at=time.monotonic(), items=items, updated_at=datetime.now().astimezone().isoformat(timespec="seconds"), error="", loading=False, provider=f"{provider_name} · Market Pulse fallback")
                succeeded = True
                break
            except Exception as fallback_exc:
                errors.append(f"{provider.__name__.replace('_fetch_', '')}: {fallback_exc}")
        if not succeeded:
            with NEWS_LOCK:
                NEWS_CACHE.update(at=time.monotonic(), error="Market Pulse and fallback news feeds failed: " + "; ".join(errors), loading=False)
    finally:
        with NEWS_REFRESH_LOCK:
            NEWS_REFRESH_RUNNING = False


def fetch_gold_news(force=False):
    """Return immediately; refresh feeds in the background and retain the last good result."""
    global NEWS_REFRESH_RUNNING
    now = time.monotonic()
    with NEWS_REFRESH_LOCK:
        if not NEWS_REFRESH_RUNNING:
            with NEWS_LOCK:
                due = force or now - NEWS_CACHE["at"] >= NEWS_TTL or not NEWS_CACHE["items"]
                if due:
                    NEWS_REFRESH_RUNNING = True
                    NEWS_CACHE["loading"] = True
                    NEWS_CACHE["error"] = ""
                    threading.Thread(target=_news_refresh_worker, daemon=True).start()
    with NEWS_LOCK:
        result = dict(NEWS_CACHE)
        result["stale"] = bool(result["items"] and (result["error"] or result["loading"]))
        result["summary"] = summarize_news_tilt(result["items"])
        parsed_base = urlsplit(MARKET_PULSE_API_BASE)
        result["market_pulse"] = {**MARKET_PULSE_STATUS, "api_base": f"{parsed_base.scheme}://{parsed_base.hostname or 'localhost'}{(':'+str(parsed_base.port)) if parsed_base.port else ''}{parsed_base.path}"}
        return result


def summarize_news_tilt(items):
    positive = sum(item["tilt"] == "Potentially supportive" for item in items)
    negative = sum(item["tilt"] == "Potentially negative" for item in items)
    mixed = sum(item["tilt"] == "Mixed / unclear" for item in items)
    total = positive + negative
    if total == 0 or positive == negative:
        label = "Mixed headline picture"
    elif positive > negative:
        label = "Supportive headlines outweigh negative"
    else:
        label = "Negative headlines outweigh supportive"
    return {"label": label, "supportive": positive, "negative": negative, "mixed": mixed, "count": len(items), "note": "Headline keyword read only; it does not measure price reaction, article context, or consensus expectations."}


def _calendar_type(name, category=""):
    text=(str(name or "")+" "+str(category or "")).casefold()
    groups=[
        ("Core CPI",("core cpi","core inflation","consumer price index excluding")),
        ("CPI",("consumer price index","cpi","inflation rate")),
        ("Core PPI",("core ppi","core producer price")),
        ("PPI",("producer price index","ppi")),
        ("NFP",("non-farm payroll","nonfarm payroll","payroll employment")),
        ("ADP Employment",("adp employment","adp nonfarm")),
        ("FOMC / Fed",("fomc","federal funds rate","fed interest rate","fed rate decision","fed chair speaks","powell speaks","federal reserve chair")),
        ("ECB",("ecb interest rate","ecb rate decision","ecb press conference","lagarde speaks")),
        ("BoE",("boe interest rate","bank of england rate","boe rate decision","bailey speaks")),
        ("BoJ",("boj interest rate","bank of japan rate","boj rate decision","ueda speaks")),
        ("RBA",("rba interest rate","reserve bank of australia rate","rba rate decision")),
        ("BoC",("boc interest rate","bank of canada rate","boc rate decision","macklem speaks")),
        ("Fed Speech",("fed speech","fed member speaks","federal reserve speech","powell")),
        ("GDP",("gross domestic product","gdp growth","gdp")),
        ("Retail Sales",("retail sales",)),
        ("Manufacturing PMI",("manufacturing pmi","ism manufacturing")),
        ("Services PMI",("services pmi","ism services","non-manufacturing pmi")),
        ("PMI",("pmi",)),
        ("Jobless Claims",("jobless claims","initial claims","continuing claims")),
        ("Consumer Confidence",("consumer confidence","consumer sentiment")),
        ("Treasury Auction",("treasury auction","bill auction","note auction","bond auction")),
        ("Employment",("employment change","unemployment rate","employment report")),
        ("Central Bank",("interest rate decision","rate decision","press conference","monetary policy")),
    ]
    for label,terms in groups:
        if any(term in text for term in terms): return label
    return str(category or "Economic event")[:80]


def _calendar_severity(name, importance):
    text=str(name or "").casefold()
    level=int(importance or 0)
    critical_terms=("cpi","consumer price","non-farm payroll","nonfarm payroll","fomc","federal funds rate","interest rate decision")
    if level>=3 and any(term in text for term in critical_terms): return "critical"
    return {0:"low",1:"low",2:"medium",3:"high"}.get(level,"medium")


def _fetch_economic_calendar():
    """Fetch Forex Factory's published weekly JSON calendar export."""
    request=urllib.request.Request(
        FOREX_FACTORY_CALENDAR_URL,
        headers={"User-Agent":"GoldTradingControlCenter/1.0 (local economic calendar)","Accept":"application/json"},
    )
    try:
        with urllib.request.urlopen(request,timeout=12) as response:
            raw=json.loads(response.read(2_000_000).decode("utf-8"))
    except HTTPError as exc:
        if exc.code==429:
            raise RuntimeError("Forex Factory temporarily rate-limited calendar requests. The app will retry automatically; use the official calendar link for immediate updates.") from exc
        raise
    if not isinstance(raw,list):
        raise ValueError("Forex Factory calendar returned an unexpected response.")
    impact_levels={"low":1,"medium":2,"high":3,"holiday":0,"non-economic":0,"":0}
    countries={"USD":"United States","GBP":"United Kingdom","EUR":"Euro Area","JPY":"Japan","CAD":"Canada","AUD":"Australia","NZD":"New Zealand","CHF":"Switzerland","CNY":"China","All":"Global"}
    events=[]
    for row in raw:
        if not isinstance(row,dict): continue
        event_name=str(row.get("title") or "").strip()
        raw_date=str(row.get("date") or "").strip()
        if not event_name or not raw_date: continue
        try:
            stamp=datetime.fromisoformat(raw_date.replace("Z","+00:00"))
            if stamp.tzinfo is None: continue
            stamp=stamp.astimezone(timezone.utc)
        except ValueError:
            continue
        currency=str(row.get("country") or "").strip()
        impact=str(row.get("impact") or "").strip()
        importance=impact_levels.get(impact.casefold(),0)
        events.append({
            "id":f"ff-{currency}-{stamp.isoformat()}-{event_name[:60]}","name":event_name[:200],
            "category":_calendar_type(event_name),"country":countries.get(currency,currency or "Global"),
            "currency":currency if currency!="All" else "","time":stamp.isoformat(),"time_estimated":False,
            "forecast":str(row.get("forecast") or ""),"previous":str(row.get("previous") or ""),"actual":str(row.get("actual") or ""),
            "importance":importance,"severity":_calendar_severity(event_name,importance),"source":"Forex Factory",
            "url":"https://www.forexfactory.com/calendar","origin":"Forex Factory"})
    events.sort(key=lambda item:item["time"])
    if not events: raise ValueError("Forex Factory returned no calendar events.")
    return events


def _calendar_refresh_worker():
    global CALENDAR_REFRESH_RUNNING
    try:
        events=_fetch_economic_calendar()
        with CALENDAR_LOCK:
            CALENDAR_CACHE.update(at=time.monotonic(),events=events,updated_at=datetime.now().astimezone().isoformat(timespec="seconds"),error="",loading=False,provider="Forex Factory")
    except Exception as exc:
        with CALENDAR_LOCK:
            CALENDAR_CACHE.update(at=time.monotonic(),error=str(exc)[:400],loading=False)
    finally:
        with CALENDAR_REFRESH_LOCK: CALENDAR_REFRESH_RUNNING=False


def fetch_economic_calendar(force=False):
    global CALENDAR_REFRESH_RUNNING
    now=time.monotonic()
    with CALENDAR_REFRESH_LOCK:
        with CALENDAR_LOCK:
            due=force or now-CALENDAR_CACHE["at"]>=CALENDAR_TTL or not CALENDAR_CACHE["events"]
        if due and not CALENDAR_REFRESH_RUNNING:
            CALENDAR_REFRESH_RUNNING=True
            with CALENDAR_LOCK: CALENDAR_CACHE["loading"]=True; CALENDAR_CACHE["error"]=""
            threading.Thread(target=_calendar_refresh_worker,daemon=True).start()
    with CALENDAR_LOCK:
        result=dict(CALENDAR_CACHE)
        result["events"]=list(result["events"])
        result["configured"]=True
        result["provider"]=result.get("provider") or "Forex Factory"
        result["stale"]=bool(result["events"] and (result["loading"] or result["error"]))
        return result


def read_trades():
    if not JOURNAL.exists():
        JOURNAL.parent.mkdir(parents=True, exist_ok=True)
        with JOURNAL.open("w", newline="", encoding="utf-8-sig") as f:
            csv.DictWriter(f, fieldnames=FIELDS).writeheader()
    with JOURNAL.open(newline="", encoding="utf-8-sig") as f:
        raw = list(csv.DictReader(f))
    rows = [{k: (v or "") for k, v in row.items()} for row in raw]
    # The EA appends a row when a position opens and another when it closes.
    # Join those lifecycle events by MT5 position identifier for one trade per row.
    positions = {r.get("Ticket Number"): r for r in rows if r.get("Open Time") and not r.get("Close Time")}
    for row in rows:
        key = row.get("Ticket Number")
        if row.get("Close Time") and key in positions:
            entry = positions[key]
            for field in ("Close Time", "Exit Price", "Profit/Loss", "Win/Loss", "Exit Reason", "Exit Comment", "Balance", "Equity"):
                if row.get(field) != "": entry[field] = row[field]
            entry["Trade Duration"] = row.get("Trade Duration", "")
    return [r for r in rows if not (r.get("Close Time") and r.get("Ticket Number") in positions)]


def parse_dt(row):
    for key in ("Close Time", "Open Time", "Date"):
        raw = row.get(key, "")
        if raw:
            for fmt in ("%Y.%m.%d %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%Y.%m.%d"):
                try:
                    return datetime.strptime(raw[:19], fmt)
                except ValueError:
                    pass
    return None


def account_history_trades(live=None, force=False):
    """Convert MT5 deal history into one row per position, cached between reads."""
    now = time.monotonic()
    with MT5_LOCK:
        if live is not None and not live.get("connected"):
            return MT5_HISTORY_CACHE["rows"] if MT5_HISTORY_CACHE["available"] else []
        if not force and now - MT5_HISTORY_CACHE["at"] < MT5_HISTORY_TTL:
            return MT5_HISTORY_CACHE["rows"]
        if mt5 is None or (live is not None and not live.get("connected")):
            return []
        try:
            deals = mt5.history_deals_get(datetime(2000, 1, 1), datetime.now())
            if deals is None:
                MT5_HISTORY_CACHE.update(at=now, rows=[], available=False, error=str(mt5.last_error()), sampled_at=datetime.now().isoformat(timespec="seconds"))
                return []
            buy_type, sell_type = mt5.DEAL_TYPE_BUY, mt5.DEAL_TYPE_SELL
            entry_in = mt5.DEAL_ENTRY_IN
            entry_out = mt5.DEAL_ENTRY_OUT
            entry_out_by = getattr(mt5, "DEAL_ENTRY_OUT_BY", 3)
            entry_inout = getattr(mt5, "DEAL_ENTRY_INOUT", 2)
            grouped = {}
            for deal in deals:
                if deal.type not in (buy_type, sell_type):
                    continue
                position_id = int(getattr(deal, "position_id", 0) or 0)
                if not position_id:
                    continue
                group = grouped.setdefault(position_id, {"entries": [], "exits": [], "deals": []})
                group["deals"].append(deal)
                entry = int(deal.entry)
                if entry in (entry_in, entry_inout):
                    group["entries"].append(deal)
                if entry in (entry_out, entry_out_by, entry_inout):
                    group["exits"].append(deal)

            currency = (live or {}).get("account", {}).get("currency", "")
            result = []
            for position_id, group in grouped.items():
                entries, exits, all_deals = group["entries"], group["exits"], group["deals"]
                if not entries and not exits:
                    continue
                entries.sort(key=lambda d: (d.time_msc, d.ticket))
                exits.sort(key=lambda d: (d.time_msc, d.ticket))
                first = entries[0] if entries else all_deals[0]
                last_exit = exits[-1] if exits else None
                entry_volume = sum(float(d.volume) for d in entries)
                exit_volume = sum(float(d.volume) for d in exits)
                closed = bool(last_exit and exit_volume + 1e-8 >= entry_volume)
                open_price = (sum(float(d.price) * float(d.volume) for d in entries) / entry_volume) if entry_volume else 0.0
                exit_price = (sum(float(d.price) * float(d.volume) for d in exits) / exit_volume) if exit_volume else 0.0
                net = sum(float(d.profit) + float(d.commission) + float(d.swap) + float(getattr(d, "fee", 0) or 0) for d in all_deals)
                side = "BUY" if first.type == buy_type else "SELL"
                opened = datetime.fromtimestamp(first.time)
                closed_at = datetime.fromtimestamp(last_exit.time) if closed and last_exit else None
                reason_codes = {getattr(mt5, "DEAL_REASON_TP", -1): "Take Profit",
                                getattr(mt5, "DEAL_REASON_SL", -2): "Stop Loss",
                                getattr(mt5, "DEAL_REASON_SO", -3): "Stop Out",
                                getattr(mt5, "DEAL_REASON_CLIENT", -4): "Manual / Client",
                                getattr(mt5, "DEAL_REASON_MOBILE", -5): "Mobile",
                                getattr(mt5, "DEAL_REASON_WEB", -6): "Web",
                                getattr(mt5, "DEAL_REASON_EXPERT", -7): "Expert Advisor"}
                reason = reason_codes.get(last_exit.reason, "Account deal") if last_exit else ""
                profit_text = f"{net:.2f}" if closed else ""
                result.append({"Trade ID": str(first.ticket), "Ticket Number": str(position_id),
                    "Date": opened.strftime("%Y.%m.%d"), "Open Time": opened.strftime("%Y.%m.%d %H:%M:%S"),
                    "Close Time": closed_at.strftime("%Y.%m.%d %H:%M:%S") if closed_at else "",
                    "Trade Duration": str(closed_at - opened) if closed_at else "",
                    "Trade Status": "Closed" if closed else "Open",
                    "Symbol": str(first.symbol), "Trade Type": side,
                    "Lot Size": f"{entry_volume:.2f}", "Entry Price": f"{open_price:.8f}" if entry_volume else "",
                    "Stop Loss": "", "Take Profit": "", "Exit Price": f"{exit_price:.8f}" if exits else "",
                    "Profit/Loss": profit_text, "Account Currency": currency, "Spread": "", "Risk %": "",
                    "Win/Loss": ("Win" if net > 0 else "Loss" if net < 0 else "Flat") if closed else "Open",
                    "Strategy Name": "MT5 account history" if int(first.magic) == 0 else f"MT5 magic {int(first.magic)}",
                    "Signal Type": "Historical deal", "Entry Reason": str(first.comment or ""),
                    "Exit Reason": reason, "Exit Comment": str(last_exit.comment or "") if last_exit else "",
                    "Screenshot": "", "R Multiple": "", "Balance": "", "Equity": ""})
            result.sort(key=lambda row: (parse_dt(row) or datetime.min, row.get("Ticket Number", "")))
            MT5_HISTORY_CACHE.update(at=time.monotonic(), rows=result, available=True, error="", sampled_at=datetime.now().isoformat(timespec="seconds"))
            return result
        except Exception as exc:
            MT5_HISTORY_CACHE.update(at=time.monotonic(), rows=[], available=False, error=str(exc), sampled_at=datetime.now().isoformat(timespec="seconds"))
            return []


def combined_trades(live=None):
    live = live if live is not None else mt5_snapshot()
    journal_rows = read_trades()
    history_rows = account_history_trades(live)
    journal_by_ticket = {str(row.get("Ticket Number", "")): row for row in journal_rows if row.get("Ticket Number")}
    merged = []
    history_tickets = set()
    for history in history_rows:
        ticket = str(history.get("Ticket Number", ""))
        journal = journal_by_ticket.get(ticket)
        if journal:
            # MT5 provides the account-wide execution totals; the EA journal
            # contributes strategy, risk, stop, and screenshot annotations.
            row = dict(history)
            for field in ("Stop Loss", "Take Profit", "Spread", "Risk %", "Strategy Name", "Signal Type",
                          "Entry Reason", "Exit Reason", "Exit Comment", "Screenshot", "R Multiple", "Balance", "Equity"):
                if journal.get(field): row[field] = journal[field]
            merged.append(row)
        else:
            merged.append(history)
        history_tickets.add(ticket)
    merged.extend(row for row in journal_rows if str(row.get("Ticket Number", "")) not in history_tickets)
    merged.sort(key=lambda row: (parse_dt(row) or datetime.min, row.get("Ticket Number", "")))
    return merged


def terminal_running():
    try:
        result = subprocess.run(["tasklist", "/FI", "IMAGENAME eq terminal64.exe", "/FO", "CSV", "/NH"], capture_output=True, text=True, timeout=2)
        return "terminal64.exe" in result.stdout.lower()
    except (OSError, subprocess.TimeoutExpired):
        return False


def _num(obj, name, default=0.0):
    try: return float(getattr(obj, name, default) or 0)
    except (TypeError, ValueError): return default


def _ollama_tags(timeout=2):
    with urllib.request.urlopen(OLLAMA_URL + "/api/tags", timeout=timeout) as response:
        return json.loads(response.read())


def ensure_ollama(force=False):
    """Keep the local Ollama API available without spawning duplicate servers."""
    global OLLAMA_LAST_START
    try:
        _ollama_tags(timeout=1)
        return True
    except Exception:
        pass
    with OLLAMA_LOCK:
        now = time.monotonic()
        if not force and now - OLLAMA_LAST_START < 20:
            return False
        OLLAMA_LAST_START = now
        candidates = [shutil.which("ollama.exe"), shutil.which("ollama")]
        local = Path(os.environ.get("LOCALAPPDATA", ""))
        program_files = Path(os.environ.get("ProgramFiles", r"C:\Program Files"))
        candidates.extend([str(local / "Programs" / "Ollama" / "ollama.exe"),
                           str(program_files / "Ollama" / "ollama.exe")])
        exe = next((Path(p) for p in candidates if p and Path(p).is_file()), None)
        if exe:
            flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
            try:
                subprocess.Popen([str(exe), "serve"], stdin=subprocess.DEVNULL,
                                 stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                                 creationflags=flags, close_fds=True)
                return True
            except OSError:
                pass
    return False


def installed_models():
    try:
        return _ollama_tags().get("models", [])
    except Exception:
        return []


def ollama_status():
    ensure_ollama()
    try:
        tags = _ollama_tags()
        models = [{"name": m.get("name", ""), "size": int(m.get("size", 0) or 0),
                   "modified_at": m.get("modified_at", "")}
                  for m in tags.get("models", []) if m.get("name")]
        installed = {m["name"] for m in models}
        selected_ready = OLLAMA_MODEL in installed or OLLAMA_MODEL + ":latest" in installed
        return {"online": True, "state": "model-ready" if selected_ready else "online-model-missing",
                "selected": OLLAMA_MODEL, "models": models, "recommended": OLLAMA_RECOMMENDED,
                "pull": dict(OLLAMA_PULL)}
    except Exception as exc:
        return {"online": False, "state": "offline", "selected": OLLAMA_MODEL,
                "models": [], "recommended": OLLAMA_RECOMMENDED,
                "pull": dict(OLLAMA_PULL), "error": str(exc)}


def save_ollama_model(model):
    global OLLAMA_MODEL
    OLLAMA_MODEL = model
    MODEL_CONFIG.write_text(json.dumps({"model": model}, indent=2), encoding="utf-8")


def _pull_model_worker(model):
    global OLLAMA_PULL
    try:
        payload = json.dumps({"model": model, "stream": True}).encode()
        req = urllib.request.Request(OLLAMA_URL + "/api/pull", data=payload,
                                    headers={"Content-Type": "application/json"})
        with urllib.request.urlopen(req, timeout=3600) as response:
            for raw in response:
                if not raw.strip(): continue
                event = json.loads(raw)
                with OLLAMA_LOCK:
                    OLLAMA_PULL.update(status=event.get("status", "Downloading"),
                                       completed=int(event.get("completed", 0) or 0),
                                       total=int(event.get("total", 0) or 0))
                if event.get("error"):
                    raise RuntimeError(event["error"])
        save_ollama_model(model)
        with OLLAMA_LOCK:
            OLLAMA_PULL.update(running=False, status="Download complete · model selected", error="")
    except Exception as exc:
        with OLLAMA_LOCK:
            OLLAMA_PULL.update(running=False, status="Download failed", error=str(exc))


def start_model_pull(model):
    global OLLAMA_PULL
    if not re.fullmatch(r"[A-Za-z0-9._:/-]{1,100}", model):
        raise ValueError("Enter a valid Ollama model tag, such as qwen3:8b.")
    ensure_ollama(force=True)
    deadline = time.monotonic() + 20
    while True:
        try:
            _ollama_tags(timeout=2)
            break
        except Exception as exc:
            if time.monotonic() >= deadline:
                raise RuntimeError("Ollama is offline. Start Ollama or check its local installation, then retry.") from exc
            time.sleep(1)
    with OLLAMA_LOCK:
        if OLLAMA_PULL.get("running"):
            raise RuntimeError("A model download is already running.")
        OLLAMA_PULL = {"running": True, "model": model, "status": "Starting download", "completed": 0, "total": 0, "error": ""}
    threading.Thread(target=_pull_model_worker, args=(model,), daemon=True).start()


def select_ollama_model(model):
    if not re.fullmatch(r"[A-Za-z0-9._:/-]{1,100}", model):
        raise ValueError("Enter a valid model name.")
    names = {m.get("name", "") for m in installed_models()}
    if model not in names and model + ":latest" not in names:
        raise LookupError("That model is not installed yet. Choose Download to get it first.")
    save_ollama_model(model)


def filter_trades(rows, params):
    def value(name): return (params.get(name) or [""])[0].strip()
    start, end = value("from"), value("to")
    symbol, side, result, query = value("symbol").casefold(), value("side").upper(), value("result").lower(), value("q").casefold()
    filtered = []
    for row in rows:
        closed = bool(row.get("Close Time"))
        if start or end:
            dt = parse_dt(row)
            if not dt: continue
            trade_day = dt.date().isoformat()
            if start and trade_day < start: continue
            if end and trade_day > end: continue
        if symbol and row.get("Symbol", "").casefold() != symbol: continue
        if side in ("BUY", "SELL") and row.get("Trade Type", "").upper() != side: continue
        pnl = number(row.get("Profit/Loss"))
        if result == "open" and closed: continue
        if result == "win" and (not closed or pnl <= 0): continue
        if result == "loss" and (not closed or pnl >= 0): continue
        if result == "flat" and (not closed or pnl != 0): continue
        if query and query not in json.dumps(row, ensure_ascii=False).casefold(): continue
        filtered.append(row)
    return filtered


def mt5_snapshot(force=False):
    """Read-only snapshot of the active MT5 terminal; this code never submits orders."""
    now = time.monotonic()
    with MT5_LOCK:
        if not force and MT5_CACHE["data"] is not None and now - MT5_CACHE["at"] < MT5_CACHE_TTL:
            return MT5_CACHE["data"]
        running = terminal_running()
        base = {"connected": False, "terminal_running": running, "state": "", "error": "",
                "account": None, "positions": [], "orders": [], "daily_pl": 0.0,
                "weekly_pl": 0.0, "monthly_pl": 0.0, "history_available": False,
                "sampled_at": datetime.now().isoformat(timespec="seconds")}
        if mt5 is None:
            base.update(state="Connector missing", error="Install the MetaTrader5 package from requirements.txt.")
        elif not running:
            base.update(state="Terminal not running", error="Start MetaTrader 5 to read live account data.")
        else:
            try:
                terminal = mt5.terminal_info()
                if not MT5_CACHE["initialized"] or terminal is None:
                    if not mt5.initialize(path=MT5_EXE, timeout=10000):
                        base.update(state="Terminal IPC unavailable", error=str(mt5.last_error()))
                        raise RuntimeError("MT5 initialize failed")
                    MT5_CACHE["initialized"] = True
                    terminal = mt5.terminal_info()
                account = mt5.account_info()
                if terminal is None or not bool(getattr(terminal, "connected", False)):
                    base.update(state="Terminal is open but disconnected", error="MT5 terminal_info.connected is false.")
                elif account is None:
                    base.update(state="No account session", error=str(mt5.last_error()))
                else:
                    mode = {getattr(mt5, "ACCOUNT_TRADE_MODE_DEMO", 0): "Demo",
                            getattr(mt5, "ACCOUNT_TRADE_MODE_CONTEST", 1): "Contest",
                            getattr(mt5, "ACCOUNT_TRADE_MODE_REAL", 2): "Real"}.get(account.trade_mode, "Unknown")
                    login = str(account.login)
                    acct = {"login_masked": ("••••" + login[-4:]) if len(login) > 4 else "••••",
                            "server": str(account.server or ""), "company": str(account.company or ""),
                            "mode": mode, "currency": str(account.currency or "USD"),
                            "leverage": int(account.leverage), "balance": _num(account, "balance"),
                            "credit": _num(account, "credit"), "equity": _num(account, "equity"),
                            "floating_profit": _num(account, "profit"), "margin": _num(account, "margin"),
                            "free_margin": _num(account, "margin_free"), "margin_level": _num(account, "margin_level"),
                            "trade_allowed": bool(account.trade_allowed), "expert_allowed": bool(account.trade_expert),
                            "terminal_trade_allowed": bool(getattr(terminal, "trade_allowed", False)),
                            "terminal_name": str(getattr(terminal, "name", "MetaTrader 5") or "MetaTrader 5"),
                            "terminal_build": int(getattr(terminal, "build", 0) or 0)}
                    raw_positions = mt5.positions_get()
                    if raw_positions is None: raise RuntimeError("positions_get failed: " + str(mt5.last_error()))
                    positions = []
                    for p in raw_positions:
                        side = "BUY" if p.type == mt5.POSITION_TYPE_BUY else "SELL"
                        action = mt5.ORDER_TYPE_BUY if side == "BUY" else mt5.ORDER_TYPE_SELL
                        volume, opened, current = _num(p, "volume"), _num(p, "price_open"), _num(p, "price_current")
                        stop, target = _num(p, "sl"), _num(p, "tp")
                        def calc_at(price_from, price_to):
                            if price_from <= 0 or price_to <= 0: return None
                            try:
                                value = mt5.order_calc_profit(action, str(p.symbol), volume, price_from, price_to)
                                return float(value) if value is not None else None
                            except Exception:
                                return None
                        stop_value, target_value = calc_at(current, stop), calc_at(current, target)
                        positions.append({"ticket": int(p.ticket), "symbol": str(p.symbol), "side": side,
                            "volume": volume, "open_price": opened, "current_price": current,
                            "stop_loss": stop, "take_profit": target, "profit": _num(p, "profit"),
                            "swap": _num(p, "swap"), "opened_at": datetime.fromtimestamp(p.time).isoformat(sep=" ", timespec="seconds"),
                            "age_seconds": max(0, int(time.time() - int(p.time))),
                            "distance_to_sl": abs(current-stop) if stop > 0 else None,
                            "distance_to_tp": abs(target-current) if target > 0 else None,
                            "remaining_risk": max(0.0, -stop_value) if stop_value is not None else None,
                            "remaining_reward": max(0.0, target_value) if target_value is not None else None,
                            "live_rr": (max(0.0, target_value)/max(0.0, -stop_value)) if target_value is not None and stop_value is not None and -stop_value > 1e-9 else None,
                            "magic": int(p.magic), "comment": str(p.comment or "")})
                    raw_orders = mt5.orders_get()
                    if raw_orders is None: raise RuntimeError("orders_get failed: " + str(mt5.last_error()))
                    orders = []
                    order_names = {getattr(mt5, "ORDER_TYPE_BUY_LIMIT", -1): "Buy Limit",
                                   getattr(mt5, "ORDER_TYPE_SELL_LIMIT", -2): "Sell Limit",
                                   getattr(mt5, "ORDER_TYPE_BUY_STOP", -3): "Buy Stop",
                                   getattr(mt5, "ORDER_TYPE_SELL_STOP", -4): "Sell Stop",
                                   getattr(mt5, "ORDER_TYPE_BUY_STOP_LIMIT", -5): "Buy Stop Limit",
                                   getattr(mt5, "ORDER_TYPE_SELL_STOP_LIMIT", -6): "Sell Stop Limit"}
                    for o in raw_orders:
                        orders.append({"ticket": int(o.ticket), "symbol": str(o.symbol),
                            "type": order_names.get(o.type, str(o.type)), "volume": _num(o, "volume_current", _num(o, "volume_initial")),
                            "price": _num(o, "price_open"), "stop_loss": _num(o, "sl"),
                            "take_profit": _num(o, "tp"), "setup_at": datetime.fromtimestamp(o.time_setup).isoformat(sep=" ", timespec="seconds")})
                    today = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
                    week_start = today - timedelta(days=today.weekday())
                    month_start = today.replace(day=1)
                    deals = mt5.history_deals_get(month_start, datetime.now())
                    daily_pl = weekly_pl = monthly_pl = 0.0
                    if deals is not None:
                        for d in deals:
                            if d.type not in (mt5.DEAL_TYPE_BUY, mt5.DEAL_TYPE_SELL): continue
                            when = datetime.fromtimestamp(d.time)
                            net = _num(d, "profit") + _num(d, "commission") + _num(d, "swap") + _num(d, "fee")
                            if when >= today: daily_pl += net
                            if when >= week_start: weekly_pl += net
                            monthly_pl += net
                    else:
                        daily_pl = weekly_pl = monthly_pl = None
                    base.update(connected=True, state="Connected", account=acct, positions=positions,
                                orders=orders, daily_pl=daily_pl, weekly_pl=weekly_pl,
                                monthly_pl=monthly_pl, history_available=deals is not None,
                                history_error="" if deals is not None else str(mt5.last_error()))
            except Exception as exc:
                if not base["error"]: base.update(state="MT5 read error", error=str(exc))
        MT5_CACHE.update(at=time.monotonic(), data=base)
        return base



def mt5_expert_logs(limit=250):
    """Return recent lines from the current MT5 MQL5/Logs expert log (read-only)."""
    log_dir = MT5_EXPERT_LOG_DIR
    if mt5 is not None:
        try:
            terminal = mt5.terminal_info()
            data_path = getattr(terminal, "data_path", "") if terminal else ""
            if data_path:
                log_dir = Path(data_path) / "MQL5" / "Logs"
        except Exception:
            pass
    try:
        files = sorted(log_dir.glob("*.log"), key=lambda item: item.stat().st_mtime, reverse=True)
        if not files:
            return {"available": False, "file": "", "lines": [], "updated_at": "", "error": "No MT5 Expert log file found. Start MT5 and attach the EA."}
        latest = files[0]
        # MT5 writes UTF-16LE logs on some terminals without a BOM. Reading
        # only the tail can also start between UTF-16 code units, so detect
        # the byte pattern and align the seek before decoding.
        file_size = latest.stat().st_size
        start = max(0, file_size - 1_000_000)
        with latest.open("rb") as stream:
            stream.seek(start)
            raw = stream.read()
        if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
            encoding = "utf-16"
        else:
            sample = raw[:4096]
            even_nuls = sample[0::2].count(0)
            odd_nuls = sample[1::2].count(0)
            if max(even_nuls, odd_nuls) > max(8, len(sample) // 40):
                # NUL placement flips when a tail read begins on an odd byte.
                even_positions_are_nuls = even_nuls > odd_nuls
                little_endian = even_positions_are_nuls == bool(start % 2)
                encoding = "utf-16-le" if little_endian else "utf-16-be"
                if start % 2:
                    raw = raw[1:]
            else:
                encoding = "utf-8-sig"
        text = raw.decode(encoding, errors="replace")
        # A tail read can begin in the middle of a log record. Drop only that
        # partial leading record; keep the final record even when it is live.
        if start > 0 and "\n" in text:
            text = text.split("\n", 1)[1]
        lines = [line for line in text.splitlines() if line.strip()][-max(1, min(int(limit), 500)):]
        entries = []
        log_date = latest.stem
        log_date = f"{log_date[0:4]}.{log_date[4:6]}.{log_date[6:8]}" if re.fullmatch(r"\d{8}", log_date) else ""
        for line in lines:
            fields = [part.strip() for part in line.split("\t")]
            if len(fields) >= 5 and re.fullmatch(r"\d{2}:\d{2}:\d{2}(?:\.\d+)?", fields[2]):
                # MT5 terminal format: event id, severity, time, source, message.
                when = f"{log_date} {fields[2]}".strip()
                event = f"ID {fields[0]} · level {fields[1]}"
                source = fields[3]
                message = "\t".join(fields[4:])
            elif len(fields) >= 3 and re.fullmatch(r"\d{4}[.\-/]\d{2}[.\-/]\d{2}\s+\d{2}:\d{2}:\d{2}(?:\.\d+)?", fields[0]):
                when, source, message = fields[0], fields[-2] if len(fields) > 3 else fields[1], "\t".join(fields[-1:])
                event = ""
            else:
                match = re.match(r"^(\d{4}[.\-/]\d{2}[.\-/]\d{2}\s+\d{2}:\d{2}:\d{2}(?:\.\d+)?)\s+(.*?)\s{2,}(.*)$", line)
                when, source, message = match.groups() if match else ("", "MT5 Expert", line.strip())
                event = ""
            entries.append({"time": when, "event": event, "source": source,
                            "message": message, "raw": line})
        return {"available": True, "file": latest.name, "lines": lines, "entries": entries,
                "encoding": encoding,
                "updated_at": datetime.fromtimestamp(latest.stat().st_mtime).isoformat(timespec="seconds"),
                "error": ""}
    except OSError as exc:
        return {"available": False, "file": "", "lines": [], "updated_at": "", "error": str(exc)}



def _ea_ladder_settings(data_dir):
    settings = {"risk_reward": 4.0, "tp_r": [1.0, 1.5, 2.0, 3.0], "partial_enabled": False}
    try:
        source = Path(data_dir) / "MQL5" / "Experts" / "AdaptiveMultiStrategyEA_REAL.mq5"
        text = source.read_text(encoding="utf-8", errors="replace")
        for name, default in (("TP1_R",1.0),("TP2_R",1.5),("TP3_R",2.0),("TP4_R",3.0),("RiskRewardRatio",4.0)):
            match = re.search(r"input\s+double\s+"+name+r"\s*=\s*([0-9]+(?:\.[0-9]+)?)", text)
            if match:
                value = float(match.group(1))
                if name == "RiskRewardRatio": settings["risk_reward"] = value
                else: settings["tp_r"][int(name[2])-1] = value
        match = re.search(r"input\s+bool\s+EnablePartialTP\s*=\s*(true|false)", text, re.I)
        if match: settings["partial_enabled"] = match.group(1).lower() == "true"
    except (OSError, ValueError):
        pass
    return settings


def _position_chart_setup(position, data_dir):
    side = "LONG" if position.type == mt5.POSITION_TYPE_BUY else "SHORT"
    buy = side == "LONG"
    entry = _num(position, "price_open")
    current_sl = _num(position, "sl")
    final_tp = _num(position, "tp")
    initial_sl = 0.0
    initial_volume = _num(position, "volume")
    position_id = int(getattr(position, "identifier", position.ticket))
    try:
        history = mt5.history_orders_get(position=position_id)
        entries = [order for order in (history or ())
                   if order.type in (mt5.ORDER_TYPE_BUY, mt5.ORDER_TYPE_SELL)
                   and _num(order, "price_open") > 0]
        entries.sort(key=lambda order: int(getattr(order, "time_setup", 0)))
        if entries:
            first = entries[0]
            if _num(first, "sl") > 0: initial_sl = _num(first, "sl")
            initial_volume = _num(first, "volume_initial", initial_volume)
    except Exception:
        pass
    config = _ea_ladder_settings(data_dir)
    risk_distance = abs(entry-initial_sl) if initial_sl > 0 else 0.0
    inferred = False
    if risk_distance <= 0 and final_tp > 0 and config["risk_reward"] > 0:
        risk_distance = abs(final_tp-entry)/config["risk_reward"]
        initial_sl = entry-risk_distance if buy else entry+risk_distance
        inferred = True
    elif risk_distance <= 0 and current_sl > 0:
        initial_sl = current_sl
        risk_distance = abs(entry-current_sl)
        inferred = True
    point = _num(mt5.symbol_info(position.symbol), "point", 0.01)
    stop_moved = current_sl > 0 and initial_sl > 0 and abs(current_sl-initial_sl) > max(point*1.5, 1e-8)
    current = _num(position, "price_current")
    levels = []
    if risk_distance > 0:
        for index, multiple in enumerate(config["tp_r"], 1):
            price = final_tp if index == 4 and final_tp > 0 else entry + (risk_distance*multiple if buy else -risk_distance*multiple)
            touched = current >= price if buy else current <= price
            levels.append({"label": f"TP{index}", "price": price, "r": multiple,
                           "touched": bool(touched), "active_ladder": bool(config["partial_enabled"])})
    return {"ticket": int(position.ticket), "side": side, "volume": _num(position, "volume"),
            "initial_volume": initial_volume, "entry": entry, "current_price": current,
            "initial_sl": initial_sl, "stop_loss": current_sl, "stop_moved": stop_moved,
            "risk_distance": risk_distance, "tp_levels": levels, "take_profit": final_tp,
            "partial_enabled": config["partial_enabled"], "risk_inferred": inferred,
            "digits": int(getattr(mt5.symbol_info(position.symbol), "digits", 2))}


def gold_chart_snapshot(timeframe="M5"):
    """Read broker-provided gold candles and current quote; never submits orders."""
    if mt5 is None:
        return {"available": False, "symbol": "", "candles": [], "error": "MetaTrader5 connector is unavailable."}
    live = mt5_snapshot()
    if not live.get("connected"):
        return {"available": False, "symbol": "", "candles": [], "error": live.get("error") or "Connect and sign in to MT5 first."}
    try:
        with MT5_LOCK:
            preferred = os.environ.get("GOLD_SYMBOL", "").strip()
            symbols = mt5.symbols_get()
            names = [str(item.name) for item in (symbols or [])]
            if preferred:
                symbol = preferred if preferred in names else ""
            else:
                candidates = [name for name in names if "XAUUSD" in name.upper() or "GOLD" in name.upper()]
                candidates.sort(key=lambda name: (0 if name.upper() == "XAUUSD" else 1 if name.upper().startswith("XAUUSD") else 2 if name.upper() == "GOLD" else 3, len(name)))
                symbol = candidates[0] if candidates else ""
            if not symbol:
                return {"available": False, "symbol": "", "candles": [], "error": "No XAUUSD/GOLD symbol found. Set GOLD_SYMBOL to your broker's exact gold symbol."}
            info = mt5.symbol_info(symbol)
            if info is None:
                return {"available": False, "symbol": symbol, "candles": [], "error": str(mt5.last_error())}
            if not getattr(info, "visible", False) and not mt5.symbol_select(symbol, True):
                return {"available": False, "symbol": symbol, "candles": [], "error": "Could not select the gold symbol in Market Watch."}
            timeframe = timeframe.upper()
            timeframe_map = {"M1": mt5.TIMEFRAME_M1, "M5": mt5.TIMEFRAME_M5,
                             "M15": mt5.TIMEFRAME_M15, "M30": mt5.TIMEFRAME_M30,
                             "H1": mt5.TIMEFRAME_H1}
            if timeframe not in timeframe_map: timeframe = "M5"
            rates = mt5.copy_rates_from_pos(symbol, timeframe_map[timeframe], 0, 600)
            if rates is None or len(rates) == 0:
                return {"available": False, "symbol": symbol, "candles": [], "error": "No M5 history returned. Open the gold chart in MT5 and wait for data."}
            candles = [{"time": int(row["time"]), "open": float(row["open"]), "high": float(row["high"]),
                        "low": float(row["low"]), "close": float(row["close"])} for row in rates]
            tick = mt5.symbol_info_tick(symbol)
            bid = float(tick.bid) if tick else candles[-1]["close"]
            digits = int(getattr(info, "digits", 2))
            raw_positions = mt5.positions_get(symbol=symbol)
            if raw_positions is None: raw_positions = ()
            terminal = mt5.terminal_info()
            data_path = getattr(terminal, "data_path", "") if terminal else ""
            data_dir = Path(data_path) if data_path else MT5_DATA_DIR
            setups = [_position_chart_setup(position, data_dir) for position in raw_positions]
            setups.sort(key=lambda setup: setup["ticket"], reverse=True)
        return {"available": True, "symbol": symbol, "candles": candles, "bid": bid, "digits": digits,
                "positions": setups, "timeframe": timeframe,
                "updated_at": datetime.now().isoformat(timespec="seconds"), "error": ""}
    except Exception as exc:
        return {"available": False, "symbol": "", "candles": [], "error": str(exc)}


def ea_status():
    try:
        age = max(0, int(time.time() - EA_HEARTBEAT.stat().st_mtime))
        text = EA_HEARTBEAT.read_text(encoding="utf-8", errors="replace").strip()
        fields = text.split("|")
        active = age <= 10 and fields and fields[0] != "STOPPED"
        return {"attached": bool(active), "age_seconds": age,
                "symbol": fields[1] if len(fields) > 1 else "",
                "trading_enabled": len(fields) > 2 and fields[2] == "ORDERS_ENABLED",
                "state": "Reporting" if active else ("Stopped" if fields and fields[0] == "STOPPED" else "No recent heartbeat")}
    except OSError:
        return {"attached": False, "age_seconds": None, "symbol": "", "trading_enabled": False, "state": "No heartbeat yet"}


def analytics(rows, live=None):
    closed = [r for r in rows if r.get("Close Time") and r.get("Profit/Loss") != ""]
    now = datetime.now()
    pl = [number(r.get("Profit/Loss")) for r in closed]
    wins = sum(p > 0 for p in pl)
    # A P/L curve starting at zero is not account equity. Anchor the
    # history to the current/last known balance before computing percentage DD.
    balance_anchor = number(live.get("account", {}).get("balance")) if live and live.get("connected") else 0.0
    if balance_anchor <= 0:
        balance_anchor = next((number(r.get("Balance")) for r in reversed(rows) if number(r.get("Balance")) > 0), 0.0)
    starting_equity = balance_anchor - sum(pl)
    drawdown_available = bool(closed and starting_equity > 0)
    equity = []
    run = starting_equity if drawdown_available else 0.0
    peak = run
    max_dd = 0.0
    for p in pl:
        run += p
        peak = max(peak, run)
        if drawdown_available and peak > 0:
            max_dd = max(max_dd, (peak - run) / peak * 100)
        equity.append(run - starting_equity if drawdown_available else run)
    today = date.today()
    start_week = today - timedelta(days=today.weekday())
    current_month = today.strftime("%Y-%m")
    daily = weekly = monthly = 0.0
    byhour, bymonth = defaultdict(float), defaultdict(float)
    for r, p in zip(closed, pl):
        dt = parse_dt(r)
        if not dt:
            continue
        if dt.date() == today: daily += p
        if dt.date() >= start_week: weekly += p
        if dt.strftime("%Y-%m") == current_month: monthly += p
        byhour[dt.hour] += p
        bymonth[dt.strftime("%Y-%m")] += p
    positive = sum(p for p in pl if p > 0)
    negative = abs(sum(p for p in pl if p < 0))
    rvalues = [number(r.get("R Multiple")) for r in closed if r.get("R Multiple") != ""]
    last = next((r for r in reversed(rows) if r.get("Balance") or r.get("Equity")), {})
    hours = sorted(byhour.items(), key=lambda x: x[1])
    return {"total_trades": len(closed), "wins": wins, "losses": len(closed)-wins,
            "win_rate": 100*wins/len(closed) if closed else 0,
            "profit_factor": positive/negative if negative else (positive if positive else 0),
            "daily_pl": (live["daily_pl"] if live and live["connected"] else daily),
            "weekly_pl": (live["weekly_pl"] if live and live["connected"] else weekly),
            "monthly_pl": (live["monthly_pl"] if live and live["connected"] else monthly),
            "drawdown": max_dd if drawdown_available else None, "max_drawdown": max_dd if drawdown_available else None, "drawdown_available": drawdown_available, "drawdown_basis": "Estimated from live balance and returned closed-trade history; deposits and withdrawals can affect the estimate." if drawdown_available else "Unavailable · a usable account-balance anchor is missing",
            "average_rr": sum(rvalues)/len(rvalues) if rvalues else 0,
            "balance": (live["account"]["balance"] if live and live["connected"] else number(last.get("Balance"))),
            "equity": (live["account"]["equity"] if live and live["connected"] else number(last.get("Equity"))),
            "floating_profit": (live["account"]["floating_profit"] if live and live["connected"] else 0.0),
            "margin": (live["account"]["margin"] if live and live["connected"] else 0.0),
            "free_margin": (live["account"]["free_margin"] if live and live["connected"] else 0.0),
            "margin_level": (live["account"]["margin_level"] if live and live["connected"] else 0.0),
            "pending_orders": (len(live["orders"]) if live and live["connected"] else 0),
            "account_currency": (live["account"]["currency"] if live and live["connected"] else (last.get("Account Currency") or "USD")),
            "open_trades": (len(live["positions"]) if live and live["connected"] else sum(not r.get("Close Time") for r in rows)),
            "best_hour": f"{max(byhour, key=byhour.get):02d}:00" if byhour else "—",
            "worst_hour": f"{min(byhour, key=byhour.get):02d}:00" if byhour else "—",
            "equity_curve": equity, "months": [{"month": k, "profit": v} for k,v in sorted(bymonth.items())]}


def trade_replay_snapshot(ticket, timeframe="M5"):
    if mt5 is None:
        return {"available":False,"error":"MetaTrader 5 history is unavailable on this computer.","candles":[]}
    live=mt5_snapshot()
    if not live.get("connected"):
        return {"available":False,"error":live.get("error") or "Connect to MT5 to load broker candles.","candles":[]}
    rows=combined_trades(live)
    trade=next((r for r in rows if str(r.get("Ticket Number"))==str(ticket)),None)
    if not trade:
        return {"available":False,"error":"That trade is not present in the available account or journal history.","candles":[]}
    def local_stamp(value):
        for fmt in ("%Y.%m.%d %H:%M:%S","%Y-%m-%d %H:%M:%S","%Y.%m.%d","%Y-%m-%d"):
            try: return datetime.strptime(str(value)[:19],fmt)
            except ValueError: pass
        return None
    opened=local_stamp(trade.get("Open Time") or trade.get("Date"))
    closed=local_stamp(trade.get("Close Time")) or datetime.now()
    if not opened or closed<=opened:
        return {"available":False,"error":"The trade needs valid entry and exit timestamps before candle replay is possible.","candles":[]}
    timeframe=str(timeframe or "M5").upper()
    timeframe_map={"M1":("TIMEFRAME_M1",60),"M5":("TIMEFRAME_M5",300),"M15":("TIMEFRAME_M15",900),"M30":("TIMEFRAME_M30",1800),"H1":("TIMEFRAME_H1",3600),"H4":("TIMEFRAME_H4",14400),"D1":("TIMEFRAME_D1",86400)}
    if timeframe not in timeframe_map: timeframe="M5"
    enum_name,seconds=timeframe_map[timeframe]
    enum=getattr(mt5,enum_name,None)
    if enum is None: return {"available":False,"error":f"The installed MT5 connector does not support {timeframe} replay.","candles":[]}
    symbol=str(trade.get("Symbol") or "")
    start=opened.astimezone(timezone.utc)
    end=closed.astimezone(timezone.utc)
    truncated=False
    maximum=3000
    if (end-start).total_seconds()/seconds>maximum:
        end=start+timedelta(seconds=seconds*maximum);truncated=True
    try:
        with MT5_LOCK:
            rates=mt5.copy_rates_range(symbol,enum,start,end)
        if rates is None or len(rates)==0:
            return {"available":False,"symbol":symbol,"error":"No broker candles are available for this trade window. Open the symbol chart in MT5 and load its history.","candles":[]}
        candles=[{"time":int(x["time"]),"open":float(x["open"]),"high":float(x["high"]),"low":float(x["low"]),"close":float(x["close"])} for x in rates]
        side=str(trade.get("Trade Type") or "BUY").upper();entry=number(trade.get("Entry Price"));exit_price=number(trade.get("Exit Price"));volume=number(trade.get("Lot Size"))
        favorable=max((max(0.0,(x["high"]-entry) if side=="BUY" else (entry-x["low"])) for x in candles),default=0.0)
        adverse=max((max(0.0,(entry-x["low"]) if side=="BUY" else (x["high"]-entry)) for x in candles),default=0.0)
        mfe_money=mae_money=None
        try:
            action=mt5.ORDER_TYPE_BUY if side=="BUY" else mt5.ORDER_TYPE_SELL
            best=entry+favorable if side=="BUY" else entry-favorable
            worst=entry-adverse if side=="BUY" else entry+adverse
            m=mt5.order_calc_profit(action,symbol,volume,entry,best);a=mt5.order_calc_profit(action,symbol,volume,entry,worst)
            mfe_money=float(m) if m is not None else None;mae_money=abs(float(a)) if a is not None else None
        except Exception: pass
        return {"available":True,"symbol":symbol,"timeframe":timeframe,"candles":candles,"trade":trade,"truncated":truncated,
            "mfe_price":favorable,"mae_price":adverse,"mfe_money":mfe_money,"mae_money":mae_money,
            "entry_price":entry,"exit_price":exit_price,"side":side,"currency":live.get("account",{}).get("currency",""),
            "source":"MetaTrader 5 historical candles","sampled_at":datetime.now().astimezone().isoformat(timespec="seconds")}
    except Exception as exc:
        return {"available":False,"symbol":symbol,"error":str(exc)[:300],"candles":[]}


def workbook(rows, live=None):
    rows = [r if isinstance(r, dict) else {} for r in (rows or [])]
    wb = Workbook()
    log = wb.active; log.title = "Trade Log"
    log.append(excel_row(FIELDS[:18] + ["Screenshot"]))
    entry = wb.create_sheet("Entry Analysis"); entry.append(excel_row(["Ticket Number", "Strategy Name", "Signal Type", "Entry Reason"]))
    exitws = wb.create_sheet("Exit Analysis"); exitws.append(excel_row(["Ticket Number", "Exit Reason", "Exit Comment"]))
    dailyws = wb.create_sheet("Daily Statistics"); dailyws.append(excel_row(["Date", "Trades Taken", "Wins", "Losses", "Total Profit", "Win Rate", "Drawdown", "Average RR"]))
    monthlyws = wb.create_sheet("Monthly Report"); monthlyws.append(excel_row(["Month", "Profit", "Trades", "Win Rate", "Drawdown", "Profit Factor"]))
    dash = wb.create_sheet("Dashboard", 0)
    for r in rows:
        log.append(excel_row([r.get(k, "") for k in FIELDS[:18]] + [r.get("Screenshot", "")]))
        entry.append(excel_row([r.get(k, "") for k in ("Ticket Number", "Strategy Name", "Signal Type", "Entry Reason")]))
        exitws.append(excel_row([r.get(k, "") for k in ("Ticket Number", "Exit Reason", "Exit Comment")]))
    daily_map, month_map = defaultdict(list), defaultdict(list)
    for r in rows:
        if not r.get("Close Time"): continue
        dt = parse_dt(r)
        if not dt: continue
        daily_map[dt.date().isoformat()].append(r); month_map[dt.strftime("%Y-%m")].append(r)
    for day, group in sorted(daily_map.items()):
        ps = [number(x.get("Profit/Loss")) for x in group]; wins = sum(x > 0 for x in ps)
        dailyws.append(excel_row([day, len(group), wins, len(group)-wins, sum(ps), wins/len(group), 0, sum(number(x.get("R Multiple")) for x in group)/len(group)]))
    for month, group in sorted(month_map.items()):
        ps = [number(x.get("Profit/Loss")) for x in group]; w=sum(x>0 for x in ps); loss=abs(sum(x for x in ps if x<0))
        monthlyws.append(excel_row([month, sum(ps), len(group), w/len(group), 0, sum(x for x in ps if x>0)/loss if loss else 0]))
    m = analytics(rows, live)
    dash.append(["XUZERA TRADING OS"]); dash.append(["Metric", "Current value"])
    for label, value in [("MT5 Connection",live["state"] if live else "Unavailable"),("Account",live["account"]["login_masked"] if live and live.get("account") else "—"),("Broker",live["account"]["server"] if live and live.get("account") else "—"),("Account Mode",live["account"]["mode"] if live and live.get("account") else "—"),("Current Balance",m["balance"]),("Current Equity",m["equity"]),("Floating P/L",m["floating_profit"]),("Margin Used",m["margin"]),("Free Margin",m["free_margin"]),("Margin Level",m["margin_level"]),("Open Positions",m["open_trades"]),("Pending Orders",m["pending_orders"]),("Win Rate",m["win_rate"]/100),("Daily P/L",m["daily_pl"]),("Weekly P/L",m["weekly_pl"]),("Monthly P/L",m["monthly_pl"]),("Current Drawdown",m["drawdown"]/100),("Profit Factor",m["profit_factor"]),("Best Trading Hour",m["best_hour"]),("Worst Trading Hour",m["worst_hour"]),("AI Status","Optional — local Ollama")]: dash.append(excel_row([label,value]))
    dash.append([]); dash.append(["Trade equity curve (cumulative realized P/L)", "P/L"])
    for i,v in enumerate(m["equity_curve"],1): dash.append([i,v])
    if m["equity_curve"]:
        chart=LineChart(); chart.title="Equity Curve"; chart.y_axis.title="Cumulative P/L"; chart.x_axis.title="Closed trade"; chart.add_data(Reference(dash,min_col=2,min_row=17,max_row=17+len(m["equity_curve"])),titles_from_data=True); chart.set_categories(Reference(dash,min_col=1,min_row=18,max_row=17+len(m["equity_curve"]))); dash.add_chart(chart,"D2")
    # Chart data lives in the dashboard's far-right columns to keep the six requested sheets tidy.
    dash.cell(1, 10, "Outcome"); dash.cell(1, 11, "Count")
    dash.cell(2, 10, "Wins"); dash.cell(2, 11, m["wins"])
    dash.cell(3, 10, "Losses"); dash.cell(3, 11, m["losses"])
    dash.cell(1, 13, "Side"); dash.cell(1, 14, "Trades")
    for i, side in enumerate(("BUY", "SELL"), 2):
        dash.cell(i, 13, side); dash.cell(i, 14, sum(r.get("Trade Type", "").upper() == side for r in rows))
    dash.cell(1, 16, "Trade"); dash.cell(1, 17, "R Multiple")
    risk_rows = [r for r in rows if r.get("R Multiple") not in (None, "")]
    for i, r in enumerate(risk_rows, 2): dash.cell(i, 16, i-1); dash.cell(i, 17, number(r.get("R Multiple")))
    if monthlyws.max_row > 1:
        monthly_chart=BarChart(); monthly_chart.title="Monthly Profit"; monthly_chart.y_axis.title="Profit"; monthly_chart.x_axis.title="Month"; monthly_chart.add_data(Reference(monthlyws,min_col=2,min_row=1,max_row=monthlyws.max_row),titles_from_data=True); monthly_chart.set_categories(Reference(monthlyws,min_col=1,min_row=2,max_row=monthlyws.max_row)); dash.add_chart(monthly_chart,"D18")
    if m["total_trades"]:
        pie=PieChart(); pie.title="Win / Loss Distribution"; pie.add_data(Reference(dash,min_col=11,min_row=1,max_row=3),titles_from_data=True); pie.set_categories(Reference(dash,min_col=10,min_row=2,max_row=3)); dash.add_chart(pie,"L18")
    if sum(dash.cell(i,14).value or 0 for i in (2,3)):
        dist=PieChart(); dist.title="Trade Distribution"; dist.add_data(Reference(dash,min_col=14,min_row=1,max_row=3),titles_from_data=True); dist.set_categories(Reference(dash,min_col=13,min_row=2,max_row=3)); dash.add_chart(dist,"D34")
    if risk_rows:
        risk_chart=BarChart(); risk_chart.title="Risk Analysis — R Multiple by Trade"; risk_chart.y_axis.title="R"; risk_chart.add_data(Reference(dash,min_col=17,min_row=1,max_row=1+len(risk_rows)),titles_from_data=True); risk_chart.set_categories(Reference(dash,min_col=16,min_row=2,max_row=1+len(risk_rows))); dash.add_chart(risk_chart,"L34")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        for c in ws[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="15243B")
        for col in ws.columns:
            letter=get_column_letter(col[0].column); width=min(38,max(12,max((len(str(c.value or "")) for c in col[:150]),default=10)+2)); ws.column_dimensions[letter].width=width
    dash["B6"].number_format="0.0%"; dash["B10"].number_format="0.00%"
    # Save beside the workbook, then replace it only after a complete successful write.
    temp_book = BOOK.with_name(BOOK.stem + ".tmp.xlsx")
    try:
        wb.save(temp_book)
        temp_book.replace(BOOK)
    finally:
        if temp_book.exists():
            temp_book.unlink()


def refresh():
    live = mt5_snapshot()
    workbook(combined_trades(live), live)


def ai_report(rows):
    m=analytics(rows)
    compact=[{k:r.get(k,"") for k in ("Close Time","Trade Type","Profit/Loss","Entry Reason","Exit Reason","R Multiple")} for r in rows if r.get("Close Time")][-300:]
    prompt=("Analyze this MT5 gold trade journal for performance. Return win rate, profit factor, drawdown, best/worst trading hours if supported, common win/loss reasons and three testable improvements. State that this is descriptive, not a guarantee. Do not recommend live trading or alter settings. Summary metrics: "+json.dumps({k:v for k,v in m.items() if k not in ("equity","equity_curve","months")})+". Trades: "+json.dumps(compact,ensure_ascii=False))
    req=urllib.request.Request(OLLAMA_URL+"/api/generate",data=json.dumps({"model":OLLAMA_MODEL,"prompt":prompt,"stream":False,"think":False,"options":{"temperature":0.3}}).encode(),headers={"Content-Type":"application/json"})
    with urllib.request.urlopen(req,timeout=180) as res:
        result=json.loads(res.read())
    report=result.get("response") or result.get("message",{}).get("content","")
    return report.strip() or "The local model returned an empty response. Try again or select a different Ollama model."


def local_ai_chat(message, history=None):
    """Answer a dashboard chat question using local Ollama and read-only trading context."""
    message = str(message or "").strip()
    if not message or len(message) > 1200:
        raise ValueError("Enter a message under 1,200 characters.")
    live = mt5_snapshot()
    rows = combined_trades(live)
    metrics = analytics(rows)
    snapshot = {
        "account": {k: live.get("account", {}).get(k) for k in ("mode", "server", "currency", "balance", "equity", "margin_free")},
        "open_positions": [{k: p.get(k) for k in ("symbol", "side", "volume", "open_price", "current_price", "stop_loss", "take_profit", "profit")} for p in live.get("positions", [])[:12]],
        "performance": {k: metrics.get(k) for k in ("total_trades", "wins", "losses", "win_rate", "profit_factor", "drawdown", "daily_pl", "weekly_pl", "monthly_pl", "average_rr")},
        "recent_closed_trades": [{k: r.get(k, "") for k in ("Close Time", "Symbol", "Trade Type", "Profit/Loss", "Entry Reason", "Exit Reason", "R Multiple")} for r in rows if r.get("Close Time")][-12:]
    }
    system = ("You are the private, local-only Xuzera assistant. Answer clearly and briefly using the supplied read-only account and journal snapshot. You may explain market/session/news context and trading terminology, but you cannot place, modify, or close trades and must never claim to do so. Do not promise profit or give certainty about price direction. For trade decisions, explain uncertainty and risk. If the snapshot lacks a detail, say so. Context: " + json.dumps(snapshot, ensure_ascii=False))
    messages = [{"role": "system", "content": system}]
    for item in (history or [])[-8:]:
        if not isinstance(item, dict) or item.get("role") not in ("user", "assistant"):
            continue
        content = str(item.get("content", ""))[:1200]
        if content.strip():
            messages.append({"role": item["role"], "content": content})
    messages.append({"role": "user", "content": message})
    request = urllib.request.Request(OLLAMA_URL + "/api/chat", data=json.dumps({"model": OLLAMA_MODEL, "messages": messages, "stream": False, "options": {"temperature": 0.35}}).encode(), headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(request, timeout=180) as response:
        result = json.loads(response.read())
    reply = result.get("message", {}).get("content", "") or result.get("response", "")
    return reply.strip() or "The local model returned an empty response. Try asking in a different way."

def process_status():
    live = mt5_snapshot()
    terminal = live["terminal_running"]
    ai = ollama_status()
    ollama = ai["state"]
    trades = combined_trades(live)
    stamp = JOURNAL.stat().st_mtime if JOURNAL.exists() else None
    journal_age = max(0, int(time.time() - stamp)) if stamp else None
    return {"terminal": terminal, "mt5_connected": live["connected"], "mt5_state": live["state"],
            "mt5_error": live["error"], "account": live["account"], "ea_source_exists": (ROOT / "mt5" / "GoldPriceActionEA.mq5").exists(),
            "positions_count": len(live["positions"]), "orders_count": len(live["orders"]),
            "journal_records": len(trades), "ea_journal_records": len(read_trades()),
            "history_imported_records": len(account_history_trades(live)), "ea": ea_status(), "journal_exists": JOURNAL.exists(), "journal_path": str(JOURNAL),
            "journal_age_seconds": journal_age,
            "journal_state": "MT5 account history + EA journal" if live["connected"] and MT5_HISTORY_CACHE["available"] else ("MT5 history unavailable · EA journal" if live["connected"] else ("EA journal only · MT5 disconnected" if JOURNAL.exists() else "Waiting for first journal entry")),
            "history_error": MT5_HISTORY_CACHE["error"],
            "workbook_exists": BOOK.exists(), "connector_available": mt5 is not None,
            "ollama": ollama, "ollama_model": OLLAMA_MODEL, "ollama_online": ai["online"],
            "ollama_models": ai["models"], "ollama_recommended": ai["recommended"],
            "ollama_pull": ai["pull"], "server": "online"}


def local_system_health():
    disk=shutil.disk_usage(ROOT)
    data={"platform":platform.platform(),"python":platform.python_version(),"disk_total_gb":round(disk.total/1_000_000_000,2),"disk_free_gb":round(disk.free/1_000_000_000,2),"cpu_percent":None,"memory_percent":None,"memory_available_gb":None,"network":"Not checked"}
    try:
        import psutil
        data["cpu_percent"]=psutil.cpu_percent(interval=0.15)
        memory=psutil.virtual_memory();data["memory_percent"]=memory.percent;data["memory_available_gb"]=round(memory.available/1_000_000_000,2)
    except (ImportError,OSError): pass
    try:
        with socket.create_connection(("1.1.1.1",53),timeout=1): data["network"]="Internet reachable"
    except OSError: data["network"]="Internet unavailable"
    started=time.monotonic();live=mt5_snapshot(force=True);latency=(time.monotonic()-started)*1000
    data.update({"mt5_connected":bool(live.get("connected")),"mt5_state":live.get("state","Unavailable"),"mt5_latency_ms":round(latency),"mt5_sampled_at":live.get("sampled_at",""),"ea":ea_status(),"ollama":ollama_status().get("state","Unavailable")})
    return data


def risk_estimate(payload):
    """Estimate stop-loss exposure with the connected broker's MT5 contract rules; never submits an order."""
    if not isinstance(payload, dict):
        raise ValueError("Risk estimate input must be an object.")
    if mt5 is None:
        raise RuntimeError("MetaTrader 5 connector is unavailable.")
    live = mt5_snapshot()
    if not live.get("connected") or not live.get("account"):
        raise RuntimeError("Connect and sign in to MetaTrader 5 before estimating risk.")
    symbol = str(payload.get("symbol", "")).strip()
    side = str(payload.get("side", "BUY")).upper()
    if not symbol or len(symbol) > 64 or side not in ("BUY", "SELL"):
        raise ValueError("Choose a valid broker symbol and BUY or SELL.")
    entry = number(payload.get("entry"), -1)
    stop = number(payload.get("stop"), -1)
    risk_pct = number(payload.get("risk_percent"), -1)
    if not (0 < entry < 1e12 and 0 < stop < 1e12):
        raise ValueError("Enter positive, valid entry and stop prices.")
    if (side == "BUY" and stop >= entry) or (side == "SELL" and stop <= entry):
        raise ValueError("For a BUY, the stop must be below entry. For a SELL, it must be above entry.")
    if not (0 < risk_pct <= 10):
        raise ValueError("Risk percentage must be greater than 0 and no more than 10%.")

    with MT5_LOCK:
        info = mt5.symbol_info(symbol)
        if info is None:
            raise ValueError("That symbol is not available in the connected broker terminal.")
        account = mt5.account_info()
        if account is None:
            raise RuntimeError("Could not read the connected MT5 account.")
        equity = _num(account, "equity")
        if equity <= 0:
            raise ValueError("The broker reported no positive account equity to use as the risk basis.")
        action = mt5.ORDER_TYPE_BUY if side == "BUY" else mt5.ORDER_TYPE_SELL
        one_lot = mt5.order_calc_profit(action, symbol, 1.0, entry, stop)
        if one_lot is None:
            raise RuntimeError(f"Broker risk calculation failed: {mt5.last_error()}")
        loss_per_lot = -float(one_lot)
        if loss_per_lot <= 0:
            raise ValueError("The broker calculation did not produce a loss at that stop. Check side, prices, and symbol settings.")
        step = _num(info, "volume_step")
        minimum = _num(info, "volume_min")
        maximum = _num(info, "volume_max")
        if step <= 0 or minimum <= 0 or maximum < minimum:
            raise RuntimeError("Broker volume limits are missing or invalid for this symbol.")
        budget = equity * risk_pct / 100.0
        raw_volume = budget / loss_per_lot
        steps = int((raw_volume + step * 1e-10) / step)
        precision = min(8, max(0, len(f"{step:.8f}".rstrip("0").split(".")[-1])))
        volume = round(min(maximum, steps * step), precision)
        if volume < minimum:
            min_loss = mt5.order_calc_profit(action, symbol, minimum, entry, stop)
            return {"available": True, "eligible": False, "symbol": symbol, "side": side,
                    "currency": str(getattr(account, "currency", "")), "equity": equity,
                    "risk_percent": risk_pct, "risk_budget": budget,
                    "loss_per_lot": loss_per_lot, "min_volume": minimum,
                    "minimum_volume_risk": abs(float(min_loss)) if min_loss is not None else None,
                    "message": "The minimum broker lot would exceed your selected risk budget. No size is suggested."}
        estimate = mt5.order_calc_profit(action, symbol, volume, entry, stop)
        if estimate is None:
            raise RuntimeError(f"Broker risk calculation failed: {mt5.last_error()}")
        return {"available": True, "eligible": True, "symbol": symbol, "side": side,
                "currency": str(getattr(account, "currency", "")), "equity": equity,
                "risk_percent": risk_pct, "risk_budget": budget, "entry": entry, "stop": stop,
                "loss_per_lot": loss_per_lot, "volume": volume, "volume_step": step,
                "estimated_stop_loss": abs(float(estimate)),
                "estimated_risk_percent": abs(float(estimate)) / equity * 100,
                "at_broker_maximum": raw_volume >= maximum,
                "message": "Estimate uses MT5 broker contract settings. Commission, swap, slippage, gaps, and execution differences are excluded."}



AUTOMATION_CONFIG_FILE = ROOT / "automation_settings.json"
AUTOMATION_SECRETS_FILE = ROOT / "automation_secrets.dpapi"
AUTOMATION_LOG_FILE = ROOT / "automation_alert_log.json"
AUTOMATION_DISCLAIMER = "Markets involve significant risk. This notification is for informational purposes only and does not constitute financial advice.\n\nNo algorithm, AI model, strategy, or trading system can guarantee profits or eliminate risk."
AUTOMATION_LOCK = threading.RLock()
AUTOMATION_CONFIG = {}
AUTOMATION_SECRETS = {}
AUTOMATION_LOG = []
AUTOMATION_NOTIFICATIONS = []
AUTOMATION_SEEN = set()
AUTOMATION_LAST_MARKET = {}
AUTOMATION_MAX_LOG = 500
AUTOMATION_EVENT_TERMS = ["cpi", "consumer price", "ppi", "producer price", "non-farm", "nonfarm", "nfp", "adp", "fomc", "federal reserve", "fed speech", "powell", "ecb", "boe", "boj", "gdp", "retail sales", "pmi", "manufacturing", "services pmi", "consumer confidence", "jobless claims", "ism", "treasury auction", "gold", "usd", "dollar"]
AUTOMATION_DEFAULTS = {"enabled":False,"timezone":"Asia/Kolkata","channels":{"email":False,"telegram":False,"discord":False,"browser":True,"in_app":True},"events":["CPI","Core CPI","PPI","Core PPI","NFP","ADP","FOMC","Fed Speeches","ECB","BOE","BOJ","GDP","Retail Sales","PMI","Consumer Confidence","Jobless Claims","Gold","USD"],"timings":[1440,720,240,60,30,15,5,0,-15,-30],"headline_alerts":False,"session_alerts":{"asian":True,"london":True,"new_york":True,"overlap":True},"thresholds":{"equity_drop_percent":2.0,"margin_level_percent":300.0,"drawdown_percent":5.0,"daily_loss_amount":0.0},"reports":{"daily":False,"daily_time":"07:00","weekly":False,"weekly_day":0,"weekly_time":"07:00","monthly":False,"monthly_day":1,"monthly_time":"07:00"},"templates":{"subject":"Xuzera · {{title}}","footer":AUTOMATION_DISCLAIMER}}

def _automation_blob(data):
    import ctypes
    from ctypes import wintypes
    class Blob(ctypes.Structure):
        _fields_=[("cbData",wintypes.DWORD),("pbData",ctypes.POINTER(ctypes.c_ubyte))]
    raw=bytes(data); buf=(ctypes.c_ubyte*len(raw)).from_buffer_copy(raw); inp=Blob(len(raw),buf); out=Blob()
    ok=ctypes.windll.crypt32.CryptProtectData(ctypes.byref(inp),"Xuzera notification credentials",None,None,None,1,ctypes.byref(out))
    if not ok: raise OSError("Windows could not protect notification credentials with DPAPI.")
    try: return ctypes.string_at(out.pbData,out.cbData)
    finally: ctypes.windll.kernel32.LocalFree(out.pbData)

def _automation_unblob(data):
    import ctypes
    from ctypes import wintypes
    class Blob(ctypes.Structure):
        _fields_=[("cbData",wintypes.DWORD),("pbData",ctypes.POINTER(ctypes.c_ubyte))]
    raw=bytes(data); buf=(ctypes.c_ubyte*len(raw)).from_buffer_copy(raw); inp=Blob(len(raw),buf); out=Blob()
    ok=ctypes.windll.crypt32.CryptUnprotectData(ctypes.byref(inp),None,None,None,None,1,ctypes.byref(out))
    if not ok: raise OSError("Windows could not unlock notification credentials for this user.")
    try: return ctypes.string_at(out.pbData,out.cbData)
    finally: ctypes.windll.kernel32.LocalFree(out.pbData)

def _automation_load():
    global AUTOMATION_CONFIG,AUTOMATION_SECRETS,AUTOMATION_LOG
    try:
        saved=json.loads(AUTOMATION_CONFIG_FILE.read_text(encoding="utf-8")); AUTOMATION_CONFIG={**AUTOMATION_DEFAULTS,**saved}
        AUTOMATION_CONFIG["channels"]={**AUTOMATION_DEFAULTS["channels"],**saved.get("channels",{})}
        AUTOMATION_CONFIG["thresholds"]={**AUTOMATION_DEFAULTS["thresholds"],**saved.get("thresholds",{})}
        AUTOMATION_CONFIG["reports"]={**AUTOMATION_DEFAULTS["reports"],**saved.get("reports",{})}
        AUTOMATION_CONFIG["templates"]={**AUTOMATION_DEFAULTS["templates"],**saved.get("templates",{})}
    except (OSError,ValueError): AUTOMATION_CONFIG=json.loads(json.dumps(AUTOMATION_DEFAULTS))
    try:
        token=AUTOMATION_SECRETS_FILE.read_bytes()
        if token.startswith(b"DPAPI1"): AUTOMATION_SECRETS=json.loads(_automation_unblob(token[6:]).decode("utf-8"))
    except Exception: AUTOMATION_SECRETS={}
    try: AUTOMATION_LOG=json.loads(AUTOMATION_LOG_FILE.read_text(encoding="utf-8"))[-AUTOMATION_MAX_LOG:]
    except (OSError,ValueError): AUTOMATION_LOG=[]
    for row in AUTOMATION_LOG:
        if row.get("dedupe_key"): AUTOMATION_SEEN.add(row["dedupe_key"])

def _automation_save_config():
    AUTOMATION_CONFIG_FILE.write_text(json.dumps(AUTOMATION_CONFIG,ensure_ascii=False,indent=2),encoding="utf-8")

def _automation_save_secrets():
    encrypted=_automation_blob(json.dumps(AUTOMATION_SECRETS,ensure_ascii=False).encode("utf-8"))
    AUTOMATION_SECRETS_FILE.write_bytes(b"DPAPI1"+encrypted)

def automation_public_config():
    c=json.loads(json.dumps(AUTOMATION_CONFIG)); c["credential_status"]={"email":bool(AUTOMATION_SECRETS.get("smtp_password")),"telegram":bool(AUTOMATION_SECRETS.get("telegram_token") and AUTOMATION_SECRETS.get("telegram_chat_id")),"discord":bool(AUTOMATION_SECRETS.get("discord_webhook")),"sender":bool(AUTOMATION_SECRETS.get("smtp_user")),"recipients":bool(AUTOMATION_SECRETS.get("recipients"))}; c["disclaimer"]=AUTOMATION_DISCLAIMER; c["sender_address"]=AUTOMATION_SECRETS.get("smtp_user",""); c["recipients"]=AUTOMATION_SECRETS.get("recipients",[]); c["smtp_host"]=AUTOMATION_SECRETS.get("smtp_host","smtp.gmail.com"); c["smtp_port"]=AUTOMATION_SECRETS.get("smtp_port",587); c["smtp_starttls"]=AUTOMATION_SECRETS.get("smtp_starttls",True); c["telegram_chat_id"]=AUTOMATION_SECRETS.get("telegram_chat_id",""); c["selected_model"]=OLLAMA_MODEL
    return c

def automation_log(title,body,kind="info",dedupe_key="",delivery=None):
    if dedupe_key and dedupe_key in AUTOMATION_SEEN: return None
    row={"id":str(__import__("uuid").uuid4()),"created_at":datetime.now().astimezone().isoformat(timespec="seconds"),"title":str(title)[:200],"body":(str(body)+"\n\n"+AUTOMATION_DISCLAIMER)[:10000],"kind":kind,"dedupe_key":dedupe_key,"delivery":delivery or {},"read":False}
    with AUTOMATION_LOCK:
        if dedupe_key: AUTOMATION_SEEN.add(dedupe_key)
        AUTOMATION_LOG.append(row); del AUTOMATION_LOG[:-AUTOMATION_MAX_LOG]
        AUTOMATION_NOTIFICATIONS.append(row); del AUTOMATION_NOTIFICATIONS[:-200]
        try: AUTOMATION_LOG_FILE.write_text(json.dumps(AUTOMATION_LOG,ensure_ascii=False,indent=2),encoding="utf-8")
        except OSError: pass
    return row

def _automation_email(subject,body):
    secret=AUTOMATION_SECRETS
    host=str(secret.get("smtp_host") or "smtp.gmail.com"); port=int(secret.get("smtp_port") or (587 if secret.get("smtp_starttls",True) else 465)); user=str(secret.get("smtp_user") or ""); password=str(secret.get("smtp_password") or ""); recipients=secret.get("recipients") or []
    if not user or not password or not recipients: raise RuntimeError("Configure sender address, app password, and at least one recipient first.")
    template=str(AUTOMATION_CONFIG.get("templates",{}).get("subject","{{title}}")); subject=template.replace("{{title}}",subject)[:200]
    msg=EmailMessage(); msg["Subject"]=subject; msg["From"]=user; msg["To"]=", ".join(recipients); msg.set_content(body)
    if secret.get("smtp_starttls",True):
        with smtplib.SMTP(host,port,timeout=20) as server: server.ehlo(); server.starttls(context=ssl.create_default_context()); server.ehlo(); server.login(user,password); server.send_message(msg)
    else:
        with smtplib.SMTP_SSL(host,port,timeout=20,context=ssl.create_default_context()) as server: server.login(user,password); server.send_message(msg)
    return f"Email sent to {len(recipients)} recipient(s)"

def _automation_telegram(body):
    token=AUTOMATION_SECRETS.get("telegram_token"); chat=AUTOMATION_SECRETS.get("telegram_chat_id")
    if not token or not chat: raise RuntimeError("Configure Telegram bot token and chat ID first.")
    text=body[:3600].rstrip()+"\n\n"+AUTOMATION_DISCLAIMER
    request=urllib.request.Request(f"https://api.telegram.org/bot{token}/sendMessage",data=json.dumps({"chat_id":chat,"text":text[:4090]}).encode(),headers={"Content-Type":"application/json"})
    with urllib.request.urlopen(request,timeout=15) as response: data=json.loads(response.read())
    if not data.get("ok"): raise RuntimeError(str(data.get("description") or "Telegram delivery failed."))
    return "Telegram sent"

def _automation_discord(body):
    webhook=AUTOMATION_SECRETS.get("discord_webhook")
    if not webhook or not webhook.startswith("https://discord.com/api/webhooks/"): raise RuntimeError("Configure a valid Discord webhook URL first.")
    text=body[:1450].rstrip()+"\n\n"+AUTOMATION_DISCLAIMER
    request=urllib.request.Request(webhook,data=json.dumps({"content":text[:1990]}).encode(),headers={"Content-Type":"application/json"})
    with urllib.request.urlopen(request,timeout=15) as response: response.read(4096)
    return "Discord sent"

def automation_deliver(title,body,kind="info",dedupe_key="",channels=None):
    c=AUTOMATION_CONFIG; selected=channels or c.get("channels",{}); statuses={}
    message=f"{title}\n\n{body}\n\n{AUTOMATION_DISCLAIMER}"
    if selected.get("email"):
        try: statuses["email"]={"ok":True,"message":_automation_email(title,message)}
        except Exception as exc: statuses["email"]={"ok":False,"error":str(exc)[:400]}
    if selected.get("telegram"):
        try: statuses["telegram"]={"ok":True,"message":_automation_telegram(message)}
        except Exception as exc: statuses["telegram"]={"ok":False,"error":str(exc)[:400]}
    if selected.get("discord"):
        try: statuses["discord"]={"ok":True,"message":_automation_discord(message)}
        except Exception as exc: statuses["discord"]={"ok":False,"error":str(exc)[:400]}
    return automation_log(title,body,kind,dedupe_key,statuses)

def _automation_event_match(event):
    text=(str(event.get("name","")+" "+str(event.get("category",""))+" "+str(event.get("currency","")))).casefold()
    selected=[str(x).casefold() for x in AUTOMATION_CONFIG.get("events",[])]
    return any(term in text for term in AUTOMATION_EVENT_TERMS if any(sel in term or term in sel for sel in selected))

def _automation_ai(event,live=None):
    try:
        positions=(live or {}).get("positions",[])
        facts={"event":event.get("name"),"currency":event.get("currency"),"time":event.get("time"),"forecast":event.get("forecast"),"previous":event.get("previous"),"actual":event.get("actual"),"severity":event.get("severity"),"account":{k:(live or {}).get("account",{}).get(k) for k in ("currency","equity","floating_profit","margin_level")},"ea_status":ea_status(),"open_positions":[{k:p.get(k) for k in ("symbol","side","volume","profit","stop_loss","take_profit")} for p in positions[:8]]}
        prompt="You are a local-only economic-news explainer. Using only the supplied facts, write: a short summary; why the event matters; conditional possible USD and gold reactions (never claim direction is certain); possible relevance to the listed open positions; and cautious, non-personalized risk considerations. Do not direct the user to buy, sell, reduce a specific amount, or change trades. State missing facts and uncertainty. Facts: "+json.dumps(facts,ensure_ascii=False)
        req=urllib.request.Request(OLLAMA_URL+"/api/generate",data=json.dumps({"model":OLLAMA_MODEL,"prompt":prompt,"stream":False,"think":False,"options":{"temperature":0.25}}).encode(),headers={"Content-Type":"application/json"})
        with urllib.request.urlopen(req,timeout=50) as response: data=json.loads(response.read())
        return str(data.get("response") or data.get("message",{}).get("content") or "").strip() or "Local model returned no analysis."
    except Exception as exc: return f"Local AI unavailable ({OLLAMA_MODEL}): {str(exc)[:180]}"

def _automation_briefing_ai(live,events,rows):
    try:
        context={"upcoming_events":[{k:e.get(k) for k in ("name","currency","time","forecast","previous","severity")} for e in events[:12]],"account":{"currency":(live.get("account") or {}).get("currency"),"equity":(live.get("account") or {}).get("equity"),"floating_profit":(live.get("account") or {}).get("floating_profit")},"positions":[{k:p.get(k) for k in ("symbol","side","volume","profit")} for p in live.get("positions",[])[:8]],"performance":{k:analytics(rows,live).get(k) for k in ("daily_pl","weekly_pl","monthly_pl","win_rate","drawdown")}}
        prompt="Prepare a concise morning/weekly trading briefing from this local read-only context. Give conditional Gold and USD context only when supported by scheduled events. Mention uncertainty, do not predict certainty, and give risk-management considerations without recommending a specific trade, exposure change, or order. List missing data rather than guessing. Context: "+json.dumps(context,ensure_ascii=False)
        req=urllib.request.Request(OLLAMA_URL+"/api/generate",data=json.dumps({"model":OLLAMA_MODEL,"prompt":prompt,"stream":False,"think":False,"options":{"temperature":0.25}}).encode(),headers={"Content-Type":"application/json"})
        with urllib.request.urlopen(req,timeout=60) as response: data=json.loads(response.read())
        return str(data.get("response") or data.get("message",{}).get("content") or "").strip() or "Local model returned no briefing."
    except Exception as exc: return f"Local AI unavailable ({OLLAMA_MODEL}): {str(exc)[:180]}"

def _automation_event_body(event,label,live=None,ai_text=None):
    try: when=datetime.fromisoformat(str(event.get("time","")).replace("Z","+00:00")).astimezone(ZoneInfo(AUTOMATION_CONFIG.get("timezone","Asia/Kolkata"))).strftime("%Y-%m-%d %H:%M:%S %Z")
    except Exception: when=str(event.get("time") or "time unavailable")
    details=f"Event: {event.get('name','Unknown')}\nCurrency: {event.get('currency') or 'Global'} · Impact: {event.get('severity') or 'Unrated'}\nLocal time: {when}\nForecast: {event.get('forecast') or 'Not published'} · Previous: {event.get('previous') or 'Not published'} · Actual: {event.get('actual') or 'Not published'}\nSource: {event.get('source') or 'Economic calendar'}"
    return f"{details}\n\n{label}\n\nAI summary / event analysis:\n{ai_text or 'Analysis is generated locally when the event alert is due; no directional outcome is guaranteed.'}\n\nRisk note: review current exposure and broker-provided protections. No specific position-size or trade instruction is generated."

def _automation_make_reports(now,live,rows):
    cfg=AUTOMATION_CONFIG.get("reports",{}); zone=ZoneInfo(AUTOMATION_CONFIG.get("timezone","Asia/Kolkata")); local=now.astimezone(zone); stamp=local.strftime("%Y-%m-%d")
    def send(kind,title,body,key):
        selected={**AUTOMATION_CONFIG.get("channels",{}),"email":True,"telegram":False,"discord":False}
        automation_deliver(title,body,kind,key,selected)
    if cfg.get("daily") and local.strftime("%H:%M")==cfg.get("daily_time","07:00") and f"report-daily-{stamp}" not in AUTOMATION_SEEN:
        def event_is_today(e):
            try: return datetime.fromisoformat(str(e.get("time","")).replace("Z","+00:00")).astimezone(zone).date()==local.date()
            except ValueError: return False
        events=[e for e in fetch_economic_calendar().get("events",[]) if _automation_event_match(e) and event_is_today(e)]
        ai_brief=_automation_briefing_ai(live,events,rows)
        body=f"Date: {stamp}\nTrading day overview: {len(events)} matching events on the calendar.\nUpcoming events: "+("; ".join(e.get("name","")+" · "+str(e.get("time","")) for e in events[:12]) or "No matching events currently listed.")+f"\nRisk score: Not modeled.\nVolatility forecast: Not available without a calibrated forecast model.\nSession recommendations: session openings are separately monitored; consult broker spread/liquidity.\n\nAI-generated local briefing (informational, conditional):\n{ai_brief}"
        send("report",f"Daily morning briefing · {stamp}",body,f"report-daily-{stamp}")
    if cfg.get("weekly") and local.weekday()==int(cfg.get("weekly_day",0)) and local.strftime("%H:%M")==cfg.get("weekly_time","07:00") and f"report-weekly-{local.strftime('%G-W%V')}" not in AUTOMATION_SEEN:
        end=datetime.combine(local.date(),datetime.min.time()); start=end-timedelta(days=7); period=[r for r in rows if r.get("Close Time") and (dt:=parse_dt(r)) and start<=dt<end]; vals=[number(r.get("Profit/Loss")) for r in period]; wins=sum(v>0 for v in vals); net=sum(vals); groups=defaultdict(float)
        for r in period: groups[str(r.get("Strategy Name") or "Unassigned")]+=number(r.get("Profit/Loss"))
        best=max(groups,key=groups.get) if groups else "Insufficient data"; worst=min(groups,key=groups.get) if groups else "Insufficient data"
        ai_review=ai_report(period) if period else "Local AI review skipped because the prior week had no closed trades."
        body=f"Previous week: {start.date()} through {(end-timedelta(days=1)).date()}\nWeekly profit: {net:.2f} {live.get('account',{}).get('currency','')}\nClosed trades: {len(vals)} · win rate: {wins/len(vals)*100 if vals else 0:.1f}% · realized drawdown metric: {analytics(period).get('drawdown',0):.2f}%\nBest EA/strategy by net result: {best} · worst: {worst}. Identity uses available broker labels and may be unassigned.\nBest/worst session: see Strategy page session analysis.\nMajor mistakes: not inferred from broker history.\n\nLocal AI improvement review (descriptive):\n{ai_review}"
        send("report",f"Weekly trading report · {local.strftime('%G-W%V')}",body,f"report-weekly-{local.strftime('%G-W%V')}")
    if cfg.get("monthly") and local.day==int(cfg.get("monthly_day",1)) and local.strftime("%H:%M")==cfg.get("monthly_time","07:00") and f"report-monthly-{local.strftime('%Y-%m')}" not in AUTOMATION_SEEN:
        end=datetime.combine(local.date().replace(day=1),datetime.min.time()); prior=end-timedelta(days=1); start=prior.replace(day=1); period=[r for r in rows if r.get("Close Time") and (dt:=parse_dt(r)) and start<=dt<end]; vals=[number(r.get("Profit/Loss")) for r in period]; wins=sum(v>0 for v in vals); tax_line=(f"\nTax source-ledger note: {len(period)} broker history rows are included in the prior-month summary; currency conversion and tax liability are not calculated." if cfg.get("tax") else "")
        body=f"Reporting month: {start.strftime('%B %Y')}\nMonthly realized result: {sum(vals):.2f} {live.get('account',{}).get('currency','')}\nClosed trades: {len(vals)} · win rate: {wins/len(vals)*100 if vals else 0:.1f}% · realized drawdown metric: {analytics(period).get('drawdown',0):.2f}%\nTax note: this is not tax advice or a filing calculation. Use the Tax & CA page source ledger and professional review."+tax_line
        send("report",f"Monthly trading report · {start.strftime('%B %Y')}",body,f"report-monthly-{local.strftime('%Y-%m')}")

def _automation_monitor_market(now,live,rows):
    if not AUTOMATION_CONFIG.get("enabled"): return
    session_cfg=AUTOMATION_CONFIG.get("session_alerts",{})
    session_defs=[("asian","Asian Session Open","Asia/Tokyo",9),("london","London Session Open","Europe/London",8),("new_york","New York Session Open","America/New_York",8)]
    for key,title,zone_name,hour in session_defs:
        try:
            z=ZoneInfo(zone_name); local=now.astimezone(z)
            if session_cfg.get(key) and local.weekday()<5 and local.hour==hour and local.minute==0:
                automation_deliver(title,f"The {title.removesuffix(' Open')} is opening. Session times use regional local time with daylight-saving changes. Check live spread and liquidity in your broker terminal.","session",f"session-{key}-{local.date()}")
        except Exception: pass
    try:
        london=now.astimezone(ZoneInfo("Europe/London")); newyork=now.astimezone(ZoneInfo("America/New_York"))
        if session_cfg.get("overlap") and london.weekday()<5 and newyork.weekday()<5 and 8<=london.hour<17 and newyork.hour==8 and newyork.minute==0:
            automation_deliver("London / New York session overlap", "London and New York sessions are both open. This is a calendar-based session notice, not a volatility measurement.","session",f"session-overlap-{london.date()}-{london.hour}")
    except Exception: pass
    if AUTOMATION_CONFIG.get("headline_alerts"):
        try:
            headlines=fetch_gold_news().get("items",[])
            for item in headlines[:5]:
                key="headline-"+str(item.get("url") or item.get("title") or "")[:180]
                if key and key not in AUTOMATION_SEEN:
                    automation_deliver("Market headline: "+str(item.get("title") or "News update"),str(item.get("reason") or "New market headline detected.")+" | Source: "+str(item.get("source") or "News feed"),"news",key)
        except Exception: pass
    calendar=fetch_economic_calendar(); selected=[e for e in calendar.get("events",[]) if _automation_event_match(e)]
    windows=[(1440,"24-hour reminder"),(720,"12-hour reminder"),(240,"4-hour reminder"),(60,"1-hour reminder"),(30,"30-minute reminder"),(15,"15-minute reminder"),(5,"5-minute reminder"),(0,"Event started"),(-15,"Event completed"),(-30,"Post-event analysis")]
    for event in selected:
        try: event_time=datetime.fromisoformat(str(event.get("time","")).replace("Z","+00:00"))
        except ValueError: continue
        minutes=(event_time-now).total_seconds()/60
        for offset,label in windows:
            if offset not in AUTOMATION_CONFIG.get("timings",[]): continue
            if abs(minutes-offset)<=0.55:
                key=f"event-{event.get('id')}-{offset}"
                if key in AUTOMATION_SEEN: continue
                ai=_automation_ai(event,live) if offset in (60,15,5,0,-15,-30) else None
                channels=AUTOMATION_CONFIG.get("channels",{})
                automation_deliver(f"{event.get('name','Economic event')} · {label}",_automation_event_body(event,label,live,ai),"news",key,channels)
    thresholds=AUTOMATION_CONFIG.get("thresholds",{}); account=live.get("account") or {}; state=AUTOMATION_LAST_MARKET
    equity=number(account.get("equity")); margin_level=number(account.get("margin_level")); dd=number(analytics(rows,live).get("drawdown")); daily_loss=number(analytics(rows,live).get("daily_pl"))
    equity_drop=bool(state.get("equity") and equity>0 and (state["equity"]-equity)/state["equity"]*100>=number(thresholds.get("equity_drop_percent"))>0)
    margin_low=bool(margin_level and margin_level<number(thresholds.get("margin_level_percent"))>0); dd_high=bool(dd>=number(thresholds.get("drawdown_percent"))>0); loss_limit=number(thresholds.get("daily_loss_amount")); if_loss=bool(loss_limit>0 and daily_loss<=-loss_limit)
    if equity_drop and not state.get("equity_drop"): automation_deliver("Equity drop alert",f"Equity changed from {state['equity']:.2f} to {equity:.2f} {account.get('currency','')}. Review the account in MT5.","risk",f"equity-drop-{datetime.now().strftime('%Y%m%d%H%M%S')}")
    if margin_low and not state.get("margin_low"): automation_deliver("Margin warning",f"Broker-reported margin level is {margin_level:.1f}%. Review the live account and broker requirements.","risk",f"margin-{datetime.now().date()}")
    if dd_high and not state.get("dd_high"): automation_deliver("Drawdown threshold alert",f"Current realized drawdown metric is {dd:.2f}%, at/above configured {thresholds.get('drawdown_percent')}% threshold.","risk",f"drawdown-{datetime.now().date()}")
    if if_loss and not state.get("daily_loss_hit"): automation_deliver("Daily loss threshold reached",f"Realized daily P/L is {daily_loss:.2f} {account.get('currency','')} and reached the configured loss amount {loss_limit:.2f}. No order action was taken.","risk",f"daily-loss-{datetime.now().date()}")
    state.update({"equity_drop":equity_drop,"margin_low":margin_low,"dd_high":dd_high,"daily_loss_hit":if_loss})
    state["equity"]=equity
    try:
        pos={str(p.get("ticket")):p for p in live.get("positions",[])}; before=state.get("positions")
        if before is not None:
            for ticket,p in pos.items():
                if ticket not in before: automation_deliver("Trade opened",f"Ticket {ticket} · {p.get('symbol')} {p.get('side')} · volume {p.get('volume')} · MT5 account snapshot.","trade",f"trade-open-{ticket}")
            for ticket,p in before.items():
                if ticket not in pos:
                    closed=next((r for r in rows if str(r.get("Ticket Number"))==ticket and r.get("Close Time")),{})
                    reason=str(closed.get("Exit Reason") or closed.get("Exit Comment") or "")
                    title="Stop loss hit" if "stop" in reason.casefold() or reason.casefold()=="sl" else ("Take profit hit" if "take profit" in reason.casefold() or reason.casefold()=="tp" else "Trade closed")
                    automation_deliver(title,f"Ticket {ticket} · {p.get('symbol')} {p.get('side')} · MT5 position closed. Broker result: {closed.get('Profit/Loss','not yet available')} · exit reason: {reason or 'not available'}.","trade",f"trade-close-{ticket}")
        state["positions"]=pos
    except Exception: pass
    ea=ea_status(); old_ea=state.get("ea_state")
    if old_ea is not None and ea.get("state")!=old_ea: automation_deliver(f"EA status changed · {ea.get('state')}",f"EA heartbeat state changed from {old_ea} to {ea.get('state')}. Check MT5 expert logs.","ea",f"ea-state-{datetime.now().strftime('%Y%m%d%H%M')}")
    state["ea_state"]=ea.get("state")

def automation_scheduler():
    while True:
        try:
            if AUTOMATION_CONFIG.get("enabled"):
                live=mt5_snapshot(); rows=combined_trades(live); now=datetime.now(timezone.utc)
                _automation_monitor_market(now,live,rows); _automation_make_reports(now,live,rows)
        except Exception as exc: print("Automation scheduler warning:",str(exc)[:300])
        time.sleep(30)

def _automation_read_payload(handler,limit=65536):
    length=min(int(handler.headers.get("Content-Length","0") or 0),limit)
    return json.loads(handler.rfile.read(length) or b"{}")

def automation_save_payload(payload):
    global AUTOMATION_CONFIG
    if not isinstance(payload,dict): raise ValueError("Automation settings must be an object.")
    cfg=json.loads(json.dumps(AUTOMATION_CONFIG));
    for key in ("enabled","timezone","events","timings","headline_alerts"): 
        if key in payload: cfg[key]=payload[key]
    cfg["session_alerts"]={**AUTOMATION_DEFAULTS.get("session_alerts",{}),**cfg.get("session_alerts",{})}
    for key in ("email","telegram","discord","browser","in_app"):
        if key in payload.get("channels",{}): cfg["channels"][key]=bool(payload["channels"][key])
    for key in ("asian","london","new_york","overlap"):
        if key in payload.get("session_alerts",{}): cfg["session_alerts"][key]=bool(payload["session_alerts"][key])
    for key in ("daily","daily_time","weekly","weekly_day","weekly_time","monthly","monthly_day","monthly_time"):
        if key in payload.get("reports",{}): cfg["reports"][key]=payload["reports"][key]
    for key in ("equity_drop_percent","margin_level_percent","drawdown_percent","daily_loss_amount"):
        if key in payload.get("thresholds",{}): cfg["thresholds"][key]=max(0,number(payload["thresholds"][key]))
    if "email_template" in payload: cfg["templates"]["subject"]=str(payload["email_template"])[:200]
    try: ZoneInfo(str(cfg.get("timezone","Asia/Kolkata")))
    except (ZoneInfoNotFoundError,ValueError): raise ValueError("Choose a valid IANA timezone such as Asia/Kolkata or Europe/London.")
    if not isinstance(cfg.get("events"),list): raise ValueError("Event selections must be a list.")
    cfg["events"]=[str(x)[:80] for x in cfg["events"][:60]]; cfg["timings"]=[int(x) for x in cfg.get("timings",[]) if str(x).lstrip("-").isdigit()][:20]
    secret_preview={**AUTOMATION_SECRETS}
    for source,dest in {"smtp_user":"smtp_user","smtp_password":"smtp_password","telegram_token":"telegram_token","telegram_chat_id":"telegram_chat_id","discord_webhook":"discord_webhook"}.items():
        if payload.get(source): secret_preview[dest]=payload[source]
    if "recipients" in payload: secret_preview["recipients"]=[x.strip() for x in str(payload["recipients"]).replace(";",",").split(",") if x.strip()]
    if cfg.get("enabled"):
        missing=[]
        if cfg["channels"].get("email") and not (secret_preview.get("smtp_user") and secret_preview.get("smtp_password") and secret_preview.get("recipients")): missing.append("Gmail sender, app password, and recipient")
        if cfg["channels"].get("telegram") and not (secret_preview.get("telegram_token") and secret_preview.get("telegram_chat_id")): missing.append("Telegram bot token and chat ID")
        if cfg["channels"].get("discord") and not secret_preview.get("discord_webhook"): missing.append("Discord webhook")
        if missing: raise ValueError("Automation cannot be enabled until these selected channels are configured: "+"; ".join(missing))
    AUTOMATION_CONFIG=cfg; _automation_save_config()
    secrets_changed=False
    secret_fields={"smtp_user":"smtp_user","smtp_password":"smtp_password","smtp_host":"smtp_host","smtp_port":"smtp_port","smtp_starttls":"smtp_starttls","telegram_token":"telegram_token","telegram_chat_id":"telegram_chat_id","discord_webhook":"discord_webhook"}
    for source,dest in secret_fields.items():
        if source in payload and payload[source] not in (None,""): AUTOMATION_SECRETS[dest]=payload[source]; secrets_changed=True
    if "recipients" in payload: AUTOMATION_SECRETS["recipients"]=[x.strip() for x in str(payload["recipients"]).replace(";",",").split(",") if x.strip()][:50]; secrets_changed=True
    for key in payload.get("clear_secrets",[]):
        if key in AUTOMATION_SECRETS: AUTOMATION_SECRETS.pop(key,None); secrets_changed=True
    if secrets_changed: _automation_save_secrets()
    return automation_public_config()

_automation_load()

class Handler(BaseHTTPRequestHandler):
    def do_GET(self):
        route = urlsplit(self.path)
        if route.path == "/Xuzera.ico" and APP_ICON.exists():
            body = APP_ICON.read_bytes(); self.send_response(200); self.send_header("Content-Type", "image/x-icon"); self.send_header("Cache-Control", "public, max-age=86400"); self.send_header("Content-Length", str(len(body))); self.end_headers(); self.wfile.write(body)
        elif route.path == "/Xuzera-icon.png" and (ROOT / "Xuzera-icon.png").exists():
            body=(ROOT / "Xuzera-icon.png").read_bytes(); self.send_response(200); self.send_header("Content-Type","image/png"); self.send_header("Cache-Control","public, max-age=86400"); self.send_header("Content-Length",str(len(body))); self.end_headers(); self.wfile.write(body)
        elif route.path == "/Xuzera-logo.jpg" and (ROOT / "Xuzera-logo.jpg").exists():
            body=(ROOT / "Xuzera-logo.jpg").read_bytes(); self.send_response(200); self.send_header("Content-Type","image/jpeg"); self.send_header("Cache-Control","public, max-age=86400"); self.send_header("Content-Length",str(len(body))); self.end_headers(); self.wfile.write(body)
        elif route.path in ("/", "/index.html", "/positions", "/markets", "/news", "/analytics", "/expenses", "/system", "/settings", "/journal", "/strategy", "/replay", "/tax", "/platform", "/compliance", "/infrastructure", "/automation"):
            page = DASHBOARD_FILE.read_text(encoding="utf-8") if DASHBOARD_FILE.exists() else PAGE
            body=page.encode(); self.send_response(200); self.send_header("Content-Type","text/html; charset=utf-8"); self.send_header("Content-Length",str(len(body))); self.end_headers(); self.wfile.write(body)
        elif route.path=="/api/data":
            live=mt5_snapshot(); rows=combined_trades(live); stats=analytics(rows,live); payload={"trades":rows,"metrics":stats,"equity":stats["equity_curve"],"months":stats["months"],"journal":str(JOURNAL),"history":{"available":MT5_HISTORY_CACHE["available"],"imported_count":len(account_history_trades(live)),"error":MT5_HISTORY_CACHE["error"],"sampled_at":MT5_HISTORY_CACHE["sampled_at"]},"mt5":{"connected":live["connected"],"state":live["state"],"error":live["error"],"history_available":live["history_available"],"history_error":live.get("history_error",""),"sampled_at":live["sampled_at"]},"account":live["account"],"positions":live["positions"],"orders":live["orders"]}; body=json.dumps(payload).encode(); self.send_response(200); self.send_header("Content-Type","application/json"); self.send_header("Content-Length",str(len(body))); self.end_headers(); self.wfile.write(body)
        elif route.path=="/api/risk/symbols":
            if mt5 is None or not mt5_snapshot().get("connected"):
                payload={"available":False,"symbols":[],"error":"Connect and sign in to MetaTrader 5 first."}
            else:
                try:
                    with MT5_LOCK:
                        symbols=mt5.symbols_get()
                    names=sorted({str(item.name) for item in (symbols or ()) if "XAU" in str(item.name).upper() or "GOLD" in str(item.name).upper()})
                    payload={"available":True,"symbols":names,"error":""}
                except Exception as exc:
                    payload={"available":False,"symbols":[],"error":str(exc)}
            body=json.dumps(payload,ensure_ascii=False).encode("utf-8"); self.send_response(200); self.send_header("Content-Type","application/json; charset=utf-8"); self.send_header("Cache-Control","no-store"); self.send_header("Content-Length",str(len(body))); self.end_headers(); self.wfile.write(body)
        elif route.path=="/api/trades/replay":
            query=parse_qs(route.query);payload=trade_replay_snapshot(query.get("ticket",[""])[0],query.get("timeframe",["M5"])[0])
            body=json.dumps(payload,ensure_ascii=False).encode("utf-8");self.send_response(200);self.send_header("Content-Type","application/json; charset=utf-8");self.send_header("Cache-Control","no-store");self.send_header("Content-Length",str(len(body)));self.end_headers();self.wfile.write(body)
        elif route.path=="/api/system/health":
            payload=local_system_health();body=json.dumps(payload,ensure_ascii=False).encode("utf-8");self.send_response(200);self.send_header("Content-Type","application/json; charset=utf-8");self.send_header("Cache-Control","no-store");self.send_header("Content-Length",str(len(body)));self.end_headers();self.wfile.write(body)
        elif route.path=="/api/mt5/expert-logs":
            body=json.dumps(mt5_expert_logs(), ensure_ascii=False).encode("utf-8"); self.send_response(200); self.send_header("Content-Type","application/json; charset=utf-8"); self.send_header("Cache-Control","no-store"); self.send_header("Content-Length",str(len(body))); self.end_headers(); self.wfile.write(body)
        elif route.path=="/api/mt5/gold-chart":
            timeframe=parse_qs(route.query).get("timeframe",["M5"])[0]
            body=json.dumps(gold_chart_snapshot(timeframe), ensure_ascii=False).encode("utf-8"); self.send_response(200); self.send_header("Content-Type","application/json; charset=utf-8"); self.send_header("Cache-Control","no-store"); self.send_header("Content-Length",str(len(body))); self.end_headers(); self.wfile.write(body)
        elif route.path=="/api/calendar":
            payload=fetch_economic_calendar(force=parse_qs(route.query).get("refresh", [""])[0] == "1")
            body=json.dumps(payload, ensure_ascii=False).encode("utf-8"); self.send_response(200); self.send_header("Content-Type","application/json; charset=utf-8"); self.send_header("Cache-Control","no-store"); self.send_header("Content-Length",str(len(body))); self.end_headers(); self.wfile.write(body)
        elif route.path=="/api/news":
            payload=fetch_gold_news(force=parse_qs(route.query).get("refresh", [""])[0] == "1")
            body=json.dumps(payload, ensure_ascii=False).encode("utf-8"); self.send_response(200); self.send_header("Content-Type","application/json; charset=utf-8"); self.send_header("Cache-Control","no-store"); self.send_header("Content-Length",str(len(body))); self.end_headers(); self.wfile.write(body)
        elif route.path=="/api/trades.csv":
            rows=filter_trades(combined_trades(),parse_qs(route.query,keep_blank_values=True))
            output=io.StringIO(newline=""); writer=csv.DictWriter(output,fieldnames=FIELDS,extrasaction="ignore"); writer.writeheader(); writer.writerows(rows)
            body=("\ufeff"+output.getvalue()).encode("utf-8"); self.send_response(200); self.send_header("Content-Type","text/csv; charset=utf-8"); self.send_header("Content-Disposition","attachment; filename=filtered_trades.csv"); self.send_header("Content-Length",str(len(body))); self.end_headers(); self.wfile.write(body)
        elif route.path=="/api/ollama/models":
            body=json.dumps(ollama_status()).encode(); self.send_response(200); self.send_header("Content-Type","application/json"); self.send_header("Cache-Control","no-store"); self.send_header("Content-Length",str(len(body))); self.end_headers(); self.wfile.write(body)
        elif route.path=="/api/ollama/progress":
            body=json.dumps(dict(OLLAMA_PULL)).encode(); self.send_response(200); self.send_header("Content-Type","application/json"); self.send_header("Cache-Control","no-store"); self.send_header("Content-Length",str(len(body))); self.end_headers(); self.wfile.write(body)
        elif route.path=="/api/export":
            refresh(); data=BOOK.read_bytes(); self.send_response(200); self.send_header("Content-Type","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"); self.send_header("Content-Disposition","attachment; filename=Xuzera-TradingOS.xlsx"); self.send_header("Content-Length",str(len(data))); self.end_headers(); self.wfile.write(data)
        elif route.path=="/api/automation/config":
            body=json.dumps(automation_public_config(),ensure_ascii=False).encode(); self.send_response(200); self.send_header("Content-Type","application/json; charset=utf-8"); self.send_header("Cache-Control","no-store"); self.send_header("Content-Length",str(len(body))); self.end_headers(); self.wfile.write(body)
        elif route.path=="/api/automation/log":
            body=json.dumps({"logs":list(reversed(AUTOMATION_LOG[-200:]))},ensure_ascii=False).encode(); self.send_response(200); self.send_header("Content-Type","application/json; charset=utf-8"); self.send_header("Cache-Control","no-store"); self.send_header("Content-Length",str(len(body))); self.end_headers(); self.wfile.write(body)
        elif route.path=="/api/automation/notifications":
            since=parse_qs(route.query).get("since",[""])[0]; rows=[x for x in AUTOMATION_NOTIFICATIONS if not since or x["created_at"]>since]; body=json.dumps({"notifications":rows[-50:]},ensure_ascii=False).encode(); self.send_response(200); self.send_header("Content-Type","application/json; charset=utf-8"); self.send_header("Cache-Control","no-store"); self.send_header("Content-Length",str(len(body))); self.end_headers(); self.wfile.write(body)
        elif route.path=="/api/status":
            body=json.dumps(process_status()).encode(); self.send_response(200); self.send_header("Content-Type","application/json"); self.send_header("Content-Length",str(len(body))); self.end_headers(); self.wfile.write(body)
        else: self.send_error(404)
    def do_POST(self):
        try:
            if self.path in ("/api/ollama/select", "/api/ollama/pull"):
                length=min(int(self.headers.get("Content-Length", "0") or 0), 8192)
                payload=json.loads(self.rfile.read(length) or b"{}")
                model=str(payload.get("model", "")).strip()
                if self.path.endswith("/select"):
                    select_ollama_model(model); result={"ok":True,"model":OLLAMA_MODEL}; status=200
                else:
                    start_model_pull(model); result={"ok":True,"model":model,"pull":dict(OLLAMA_PULL)}; status=202
            elif self.path=="/api/risk/estimate":
                length=min(int(self.headers.get("Content-Length", "0") or 0), 8192)
                payload=json.loads(self.rfile.read(length) or b"{}")
                result=risk_estimate(payload); status=200
            elif self.path=="/api/assistant/chat":
                length=min(int(self.headers.get("Content-Length", "0") or 0), 8192)
                payload=json.loads(self.rfile.read(length) or b"{}")
                ensure_ollama()
                result={"reply":local_ai_chat(payload.get("message", ""), payload.get("history", [])), "model":OLLAMA_MODEL}; status=200
            elif self.path=="/api/analyze":
                ensure_ollama()
                result={"report":ai_report(combined_trades()),"model":OLLAMA_MODEL}; status=200
            else: self.send_error(404); return
        except LookupError as e: result={"error":str(e)}; status=409
        except (ValueError,RuntimeError) as e: result={"error":str(e)}; status=400
        except (URLError,TimeoutError,ValueError) as e: result={"error":f"Local model unavailable ({OLLAMA_MODEL}). Check Ollama and the selected model. Details: {e}"}; status=503
        body=json.dumps(result).encode(); self.send_response(status); self.send_header("Content-Type","application/json"); self.send_header("Content-Length",str(len(body))); self.end_headers(); self.wfile.write(body)
    def log_message(self, fmt,*args): print(f"[{self.log_date_time_string()}] {fmt % args}")


def watcher():
    previous=None; last_refresh=0.0; previous_history=None
    while True:
        try:
            stamp=(JOURNAL.stat().st_mtime_ns,JOURNAL.stat().st_size) if JOURNAL.exists() else None
            live=mt5_snapshot()
            rows=combined_trades(live)
            history_signature=tuple((r.get("Ticket Number"),r.get("Open Time"),r.get("Close Time"),r.get("Profit/Loss")) for r in rows)
            if stamp!=previous or history_signature!=previous_history or time.monotonic()-last_refresh>=30:
                workbook(rows,live); previous=stamp; previous_history=history_signature; last_refresh=time.monotonic()
        except Exception as exc: print("Workbook refresh failed:",exc)
        time.sleep(1)


if __name__=="__main__":
    SCREENSHOTS.mkdir(exist_ok=True)
    read_trades(); refresh()
    ensure_ollama(force=True)
    threading.Thread(target=watcher,daemon=True).start()
    threading.Thread(target=automation_scheduler,daemon=True).start()
    print(f"Xuzera: http://{HOST}:{PORT}")
    print(f"Watching journal: {JOURNAL}")
    try:
        ThreadingHTTPServer((HOST,PORT),Handler).serve_forever()
    except OSError as exc:
        if getattr(exc, "winerror", None) == 10048 or getattr(exc, "errno", None) in (48, 98, 10048):
            raise SystemExit(f"Port {PORT} is already in use. Close the other dashboard window or set a different PORT.")
        raise


